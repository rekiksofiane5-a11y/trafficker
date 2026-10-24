
#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
cm360_trafficker.py
===================
Génère le trafficking d'une campagne ad serving à partir d'un plan média
structuré (CSV ou XLSX), via une architecture d'adaptateurs :

    --platform cm360    Campaign Manager 360 (API v5, compte de service Google)
    --platform adform   Adform (API api.adform.com, OAuth 2.0 client credentials)
    --platform mock     Mode démo : simulation locale, aucun accès requis

Chaîne de création (commune à toutes les plateformes) :
    Annonceur -> Landing page -> Campagne -> Emplacements -> Créatives
    -> Annonces (rotation + affectation des placements)

Chaque adaptateur implémente l'interface `TraffickingAdapter` et gère
l'idempotence (réutilisation des objets existants par nom).

Prérequis :
    pip install google-api-python-client google-auth     # plateforme cm360
    pip install openpyxl                                 # fichiers .xlsx

Utilisation :
    python cm360_trafficker.py plan_media.csv --platform mock
    python cm360_trafficker.py plan_media.csv --dry-run
    python cm360_trafficker.py --platform cm360 --check --service-account cle.json
    python cm360_trafficker.py "https://docs.google.com/spreadsheets/d/ID" \
        --platform mock            # plan depuis un Google Sheet partagé en lecture
    python cm360_trafficker.py plan_media.csv --platform cm360 \
        --service-account cle.json --report rapport.json \
        --creatives-dir ./creas    # créatives IMAGE en chemins relatifs
    python cm360_trafficker.py --platform adform --check \
        --adform-client-id XXX --adform-client-secret YYY

Colonnes du plan média (en-têtes insensibles à la casse / accents / espaces) :
    advertiser_id          ID de l'annonceur existant (ou advertiser_name)
    advertiser_name         Nom de l'annonceur (créé s'il n'existe pas)
    campaign_name          Nom de la campagne (obligatoire)
    campaign_start        Date de début YYYY-MM-DD (obligatoire)
    campaign_end          Date de fin YYYY-MM-DD (obligatoire)
    placement_name        Nom de l'emplacement (obligatoire)
    site_id               ID du site (recommandé) ou site_name
    site_name             Nom du site (recherche, plus lent)
    size                  Format : 300x250 (obligatoire)
    pricing_type          CPM (défaut) | CPC | CPA | FLAT
    ad_name               Nom de l'annonce (obligatoire)
    creative_name         Nom de la créative (obligatoire)
    creative_type         (optionnel) REDIRECT | IMAGE | VIDEO | TRACKING — déduit si vide
    creative_category     (optionnel) display | video | tracking — force la catégorie
    snippet               Tag tiers complet (créative REDIRECT)
    creative_path         Fichier créatif partagé : image -> display, vidéo -> in-stream

    Règle de typage des créatives (si creative_type vide) :
        tag tiers (snippet)            -> REDIRECT  (display, tag tiers)
        fichier image (png/jpg/gif)    -> IMAGE     (display)
        fichier vidéo (mp4/mov/...)    -> VIDEO     (in-stream video)
        rien de partagé                -> TRACKING  (créative de tracking générée
                                                     automatiquement)
    landing_page_url      URL de destination (obligatoire)
    active                1/oui/true pour activer (défaut : 1)

Protection des données (RGPD) :
    - Minimisation : aucune donnée personnelle n'est nécessaire ; n'incluez
      pas de données de contact (emails, noms) dans le plan média.
    - Purge : les fichiers temporaires (export CSV d'un Google Sheet) sont
      supprimés immédiatement après lecture.
    - Les données du plan ne sont transmises qu'à la plateforme cible
      (CM360 ou Adform) ; sous-traitants détaillés dans le document
      « Conformité RGPD ».
"""

import argparse
import csv
import json
import os
import re
import sys
import time
import unicodedata
from dataclasses import dataclass, field
from datetime import date
from typing import Callable, Dict, List, Optional, Tuple

# ---------------------------------------------------------------------------
# Constantes
# ---------------------------------------------------------------------------

CM360_SCOPES = ["https://www.googleapis.com/auth/dfatrafficking"]

ADFORM_TOKEN_URL = "https://id.adform.com/sts/connect/token"
ADFORM_API_BASE = "https://api.adform.com"
# Scope par défaut : à ajuster selon les scopes autorisés de votre OAuth client
ADFORM_SCOPE = "https://api.adform.com/scope"

PRICING_TYPES = {
    "cpm": "PRICING_TYPE_CPM",
    "cpc": "PRICING_TYPE_CPC",
    "cpa": "PRICING_TYPE_CPA",
    "flat": "PRICING_TYPE_FLAT_RATE",
}

PLACEMENT_CATEGORIES = {"display", "video", "tracking", "audio"}

TRUE_VALUES = {"1", "oui", "yes", "true", "vrai", "o", "y"}

# ---------------------------------------------------------------------------
# 1. Lecture et normalisation du plan média
# ---------------------------------------------------------------------------

def normalize_header(h: object) -> str:
    """'Taille (px)' -> 'taille_px', 'Date de début' -> 'date_de_debut'."""
    if h is None:
        return ""
    s = str(h).strip()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^a-zA-Z0-9]+", "_", s).strip("_").lower()
    return s


def load_rows_from_path(path: str, campaign_name: str = "",
                        advertiser_name: str = "") -> List[Dict[str, str]]:
    """Charge un CSV, XLSX ou XLS et renvoie une liste de dicts {colonne: valeur}.
    Un export « Campaign spreadsheet » CM360 est reconnu à ses en-têtes et converti
    en plan (campaign_name / advertiser_name : voir cm360_bulk.bulk_to_plan_rows)."""
    import cm360_bulk
    import ingest
    import plan_agent
    grids = plan_agent.read_grids(path)
    if not any(any(any(c for c in r) for r in g) for _, g in grids):
        sys.exit("Le fichier est vide.")
    # Trafficking sheet d'agence (nomenclature et listes dans le classeur) : lecture dédiée
    import agency_sheet
    conv = agency_sheet.read_convention(path)
    if conv:
        out, report = agency_sheet.build_rows(path, grids, conv)
        if out:
            res = ingest.IngestedRows(out)
            res.ingest_report = report
            res.nomenclature = agency_sheet.nomenclature_json(conv)
            return res
    # Cases surlignées en jaune : l'agent ne lit que celles-ci
    highlights = plan_agent.read_highlights(path)
    if highlights:
        out, report = plan_agent.extract(grids, campaign_name, advertiser_name, highlights,
                                         source_name=path)
        res = ingest.IngestedRows(out)
        res.ingest_report = report
        return res
    first = grids[0][1]
    header = [normalize_header(h) for h in (first[0] if first else [])]
    rows = [dict(zip(header, r + [""] * (len(header) - len(r))))
            for r in first[1:] if any(c for c in r)]
    if cm360_bulk.is_bulk_header(header):
        return cm360_bulk.bulk_to_plan_rows(rows, campaign_name, advertiser_name)
    if plan_agent.is_canonical(grids):
        return ingest.ingest_rows(rows, campaign_name, advertiser_name)
    # Plan « tel qu'il arrive » : l'agent retrouve chaque information
    out, report = plan_agent.extract(grids, campaign_name, advertiser_name, source_name=path)
    res = ingest.IngestedRows(out)
    res.ingest_report = report
    return res


def _read_rows(path: str) -> List[Dict[str, str]]:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        try:
            import xlrd
        except ImportError:
            sys.exit("Lecture XLS impossible : installez xlrd (pip install xlrd)")
        sh = xlrd.open_workbook(path).sheet_by_index(0)
        if sh.nrows == 0:
            sys.exit("Le fichier XLS est vide.")
        header = [normalize_header(sh.cell_value(0, c)) for c in range(sh.ncols)]
        out = []
        for r in range(1, sh.nrows):
            vals = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    v = xlrd.xldate_as_datetime(cell.value, sh.book.datemode).date().isoformat()
                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                    v = str(int(cell.value))
                else:
                    v = str(cell.value).strip()
                vals.append(v)
            if any(vals):
                out.append(dict(zip(header, vals)))
        return out
    if ext in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("Lecture XLSX impossible : installez openpyxl (pip install openpyxl)")
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rows:
            sys.exit("Le fichier XLSX est vide.")
        header = [normalize_header(h) for h in rows[0]]
        out = []
        for r in rows[1:]:
            if any(v is not None and str(v).strip() != "" for v in r):
                out.append({header[i]: ("" if i >= len(r) or r[i] is None else str(r[i]).strip())
                            for i in range(len(header))})
        return out
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames:
            reader.fieldnames = [normalize_header(h) for h in reader.fieldnames]
        out = []
        for r in reader:
            clean = {k: (v.strip() if isinstance(v, str) else v) for k, v in r.items()}
            if any(v not in (None, "") for v in clean.values()):
                out.append(clean)
        return out


# ---------------------------------------------------------------------------
# 1bis. Résolution de la source du plan (fichier local / Google Sheets)
#       et des chemins de créatives
# ---------------------------------------------------------------------------

def _google_sheet_csv_url(source: str) -> str:
    """Construit l'URL d'export CSV d'un Google Sheet à partir d'une URL
    complète (docs.google.com/spreadsheets/d/...) ou d'un ID brut.
    Le sheet doit être partagé en lecture ('Tout le monde avec le lien')."""
    m = re.search(r"/spreadsheets/d/([a-zA-Z0-9_-]+)", source)
    sid = m.group(1) if m else source.strip()
    if not re.match(r"^[a-zA-Z0-9_-]{20,}$", sid):
        raise ValueError(f"URL / ID Google Sheets non reconnu : {source}")
    return f"https://docs.google.com/spreadsheets/d/{sid}/export?format=csv"


def resolve_plan_source(source: str) -> str:
    """Résout la source du plan : fichier local, URL Google Sheets ou ID de
    sheet. Télécharge l'export CSV dans un fichier temporaire si nécessaire."""
    s = (source or "").strip()
    if not s:
        raise ValueError("Source du plan manquante.")
    is_sheet = ("docs.google.com/spreadsheets" in s
                or re.match(r"^[a-zA-Z0-9_-]{20,}$", s))
    if not is_sheet:
        return s
    import tempfile
    import urllib.request
    url = _google_sheet_csv_url(s)
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            data = resp.read()
    except Exception as e:
        sys.exit(f"Téléchargement du Google Sheet impossible : {e}\n"
                 "Vérifiez le partage : 'Tout le monde avec le lien' en lecture.")
    fd, tmp = tempfile.mkstemp(suffix=".csv")
    with os.fdopen(fd, "wb") as f:
        f.write(data)
    print(f"Plan téléchargé depuis Google Sheets -> {tmp}")
    return tmp


def load_rows_from_source(source: str, campaign_name: str = "",
                          advertiser_name: str = "") -> List[Dict[str, str]]:
    """Charge les lignes du plan depuis un fichier local CSV/XLSX/XLS,
    une URL Google Sheets ou un ID de sheet."""
    return load_rows_from_path(resolve_plan_source(source), campaign_name, advertiser_name)


def resolve_creative_path(path: str, creatives_dir: str = "") -> str:
    """Résout le chemin d'une créative IMAGE : direct si valide, sinon
    relatif au dossier de créatives (--creatives-dir, ou ZIP de l'UI)."""
    if not path:
        return path
    if os.path.isabs(path) or os.path.isfile(path):
        return path
    if creatives_dir:
        return os.path.join(creatives_dir, path)
    return path


# ---------------------------------------------------------------------------
# 2. Modèle de données du plan (commun à toutes les plateformes)
# ---------------------------------------------------------------------------

@dataclass
class CreativeSpec:
    name: str
    ctype: str                      # REDIRECT | IMAGE | VIDEO | TRACKING
    auto: bool = False              # TRACKING générée automatiquement (rien de partagé)
    snippet: str = ""
    path: str = ""
    landing_page_url: str = ""
    active: bool = True
    row_index: int = 0
    message: str = ""                # message créatif (nomenclature)
    rows: List[int] = field(default_factory=list)   # lignes du plan portant cette créative
    final_name: str = ""             # nom imposé à la main (aperçu / tableau)

@dataclass
class PlacementSpec:
    name: str
    site_id: Optional[int] = None
    site_name: str = ""
    width: int = 0
    height: int = 0
    pricing_type: str = "PRICING_TYPE_CPM"
    row_index: int = 0
    category: str = ""       # display (formats IAB) | video (in-stream) | tracking (1x1)
    objective: str = ""      # objectif (notoriété, trafic...) — nomenclature
    audience: str = ""       # audience / cible — nomenclature
    targeting: str = ""      # ciblage — nomenclature
    format_label: str = ""   # format tel qu'écrit dans le plan (« Pavé »...), sinon la taille
    package: str = ""        # package CM360 (groupe d'emplacements), colonne optionnelle
    device: str = ""         # devices (desktop, mobile, multidevice...) — granularité
    rows: List[int] = field(default_factory=list)   # lignes du plan de cet emplacement
    final_name: str = ""     # nom imposé à la main (aperçu / tableau), prioritaire
    budget: str = ""         # budget de la ligne du plan (line item Adform)
    volume: str = ""         # volume d'impressions de la ligne
    extra: Dict[str, str] = field(default_factory=dict)   # jetons de nomenclature d'agence (partner, buy…)

@dataclass
class CampaignSpec:
    name: str
    advertiser_id: Optional[int] = None
    advertiser_name: str = ""
    start: str = ""
    end: str = ""
    budget_eur: Optional[float] = None      # colonne optionnelle budget_eur
    # Informations absentes du plan, laissées vides : {champ: [objets concernés]}
    missing: Dict[str, List[str]] = field(default_factory=dict)
    placements: Dict[str, PlacementSpec] = field(default_factory=dict)
    ads: Dict[str, dict] = field(default_factory=dict)   # ad_name -> {"placements": [...], "creatives": [...]}


IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".wbmp"}
HTML5_EXT = {".zip", ".adz"}
AUDIO_EXT = {".mp3", ".oga", ".ogg", ".ogv"}
VIDEO_EXT = {".mp4", ".mov", ".m4v", ".webm", ".mpg", ".mpeg", ".avi", ".wmv"}
CTYPE_ALIASES = {
    "REDIRECT": "REDIRECT", "TAG": "REDIRECT", "IMAGE": "IMAGE", "DISPLAY": "IMAGE",
    "VIDEO": "VIDEO", "INSTREAM": "VIDEO", "INSTREAM_VIDEO": "VIDEO",
    "TRACKING": "TRACKING", "HTML5": "HTML5", "HTML": "HTML5",
    "AUDIO": "AUDIO", "INSTREAM_AUDIO": "AUDIO",
}
CATEGORY_TO_CTYPE = {"display": "IMAGE", "video": "VIDEO", "tracking": "TRACKING", "audio": "AUDIO"}


def resolve_creative_type(explicit: str, category: str, snippet: str,
                          path: str) -> Tuple[str, bool, str]:
    """Détermine le type de créative. Renvoie (type, auto, erreur).
    Priorité : creative_type explicite > creative_category > déduction
    (snippet -> REDIRECT ; fichier image -> IMAGE ; fichier vidéo -> VIDEO ;
    rien de partagé -> TRACKING généré automatiquement)."""
    if explicit:
        ctype = CTYPE_ALIASES.get(explicit.strip().upper())
        if not ctype:
            return "", False, (f"creative_type inconnu '{explicit}' "
                               "(REDIRECT, IMAGE, VIDEO ou TRACKING).")
        return ctype, False, ""
    cat = (category or "").strip().lower()
    if cat:
        if cat not in CATEGORY_TO_CTYPE:
            return "", False, f"creative_category inconnue '{category}' (display, video ou tracking)."
        if cat == "display" and snippet and not path:
            return "REDIRECT", False, ""
        if cat == "display" and os.path.splitext(path)[1].lower() in HTML5_EXT:
            return "HTML5", False, ""
        return CATEGORY_TO_CTYPE[cat], False, ""
    if snippet:
        return "REDIRECT", False, ""
    if path:
        ext = os.path.splitext(path)[1].lower()
        if ext in VIDEO_EXT:
            return "VIDEO", False, ""
        if ext in IMAGE_EXT:
            return "IMAGE", False, ""
        if ext in HTML5_EXT:
            return "HTML5", False, ""
        if ext in AUDIO_EXT:
            return "AUDIO", False, ""
        return "", False, (f"format de fichier non géré '{ext or path}' "
                           "(specs CM360 — image : gif/jpg/png/wbmp ; HTML5 : zip/adz ; "
                           "vidéo : mp4/mov/m4v/avi/mpeg/mpg/webm/wmv ; audio : mp3/ogg/oga).")
    return "TRACKING", True, ""


def _ctype_label(c: "CreativeSpec") -> str:
    """Libellé lisible : « VIDEO in-stream », « TRACKING auto »..."""
    return {"VIDEO": "VIDEO in-stream", "IMAGE": "IMAGE display", "HTML5": "HTML5 display",
            "AUDIO": "AUDIO in-stream",
            "REDIRECT": "REDIRECT display"}.get(c.ctype, c.ctype) + (" auto" if c.auto else "")


def _as_bool(v: str) -> bool:
    return str(v).strip().lower() in TRUE_VALUES


GAP_LABELS = {"campaign_start": "Date de début", "campaign_end": "Date de fin",
              "landing_page_url": "URL de destination"}


def plan_from_rows(rows: List[Dict[str, str]], allow_incomplete: bool = False
                   ) -> Tuple[Dict[str, CampaignSpec], List[str]]:
    """Groupe les lignes en campagnes -> emplacements / annonces / créas.
    Renvoie (plan, erreurs). Une ligne = une créative dans une annonce.
    allow_incomplete : dates et URL absentes ne sont pas des erreurs ; elles restent
    vides et sont listées « à compléter » (camp.missing, plan_gaps). Une valeur
    présente mais invalide reste une erreur."""
    plan: Dict[str, CampaignSpec] = {}
    errors: List[str] = []
    seen_rows: set = set()

    def fail(idx: int, msg: str):
        errors.append(f"Ligne {idx + 2} : {msg}")

    def grouped(errs: List[str]) -> List[str]:
        """« Ligne 2 : X », « Ligne 3 : X »… -> « Lignes 2 à 3 : X » (une erreur par cause)."""
        order, lines = [], {}
        for e in errs:
            m = re.match(r"Ligne (\d+) : (.*)$", e, re.S)
            key = m.group(2) if m else e
            if key not in lines:
                order.append(key)
                lines[key] = []
            if m:
                lines[key].append(int(m.group(1)))
        out = []
        for key in order:
            ns = lines[key]
            hint = (" — renseignez la campagne (colonne, case jaune ou champ « Campagne »)."
                    if key.startswith("campaign_name manquant") else "")
            if not ns:
                out.append(key)
            elif len(ns) == 1:
                out.append(f"Ligne {ns[0]} : {key}{hint}")
            elif ns == list(range(ns[0], ns[-1] + 1)):
                out.append(f"Lignes {ns[0]} à {ns[-1]} : {key}{hint}")
            else:
                out.append(f"Lignes {', '.join(map(str, ns))} : {key}{hint}")
        return out

    for idx, r in enumerate(rows):
        get = lambda k, r=r: (r.get(k) or "").strip()

        campaign_name = get("campaign_name")
        placement_name = get("placement_name")
        ad_name = get("ad_name")
        creative_name = get("creative_name")
        size = get("size")
        start, end = get("campaign_start"), get("campaign_end")
        lp = get("landing_page_url")

        if not campaign_name:
            fail(idx, "campaign_name manquant.")
            continue
        if not placement_name:
            fail(idx, "placement_name manquant.")
            continue
        if not ad_name:
            fail(idx, "ad_name manquant.")
            continue
        if not creative_name:
            fail(idx, "creative_name manquant.")
            continue
        if not re.match(r"^\d{1,4}\s*[xX*]\s*\d{1,4}$", size or ""):
            fail(idx, f"size invalide ou manquant ('{size}', attendu '300x250').")
            continue
        dates_ok = True
        row_gaps: List[Tuple[str, str]] = []          # (champ, précision)
        values = {"campaign_start": start, "campaign_end": end}
        for d, label in ((start, "campaign_start"), (end, "campaign_end")):
            if not d and allow_incomplete:
                row_gaps.append((label, ""))
                dates_ok = False
                continue
            valid = bool(re.match(r"^\d{4}-\d{2}-\d{2}$", d or ""))
            if valid:
                try:
                    date.fromisoformat(d)
                except ValueError:
                    valid = False
            if not valid and allow_incomplete:       # illisible : laissée vide, à compléter
                row_gaps.append((label, f"valeur illisible « {d} »"))
                values[label] = ""
                dates_ok = False
            elif not valid:
                fail(idx, f"{label} manquant ou invalide (YYYY-MM-DD attendu)."
                     if not re.match(r"^\d{4}-\d{2}-\d{2}$", d or "")
                     else f"{label} n'est pas une date réelle ('{d}').")
                dates_ok = False
        start, end = values["campaign_start"], values["campaign_end"]
        if dates_ok and end < start:
            if allow_incomplete:
                row_gaps += [("campaign_start", f"fin {end} avant début {start}"),
                             ("campaign_end", f"fin {end} avant début {start}")]
                start = end = ""
            else:
                fail(idx, f"campaign_end ({end}) est avant campaign_start ({start}).")
        if not lp:
            if allow_incomplete:
                row_gaps.append(("landing_page_url", ""))
            else:
                fail(idx, "landing_page_url manquant.")
                continue

        w, h = [int(x) for x in re.split(r"[xX*]", size)]
        adv_id = get("advertiser_id") or None
        key = f"{(adv_id or get('advertiser_name')).lower()}|{campaign_name.lower()}"

        camp = plan.setdefault(key, CampaignSpec(
            name=campaign_name,
            advertiser_id=int(adv_id) if adv_id and adv_id.isdigit() else None,
            advertiser_name=get("advertiser_name"),
            start=start, end=end,
        ))
        if allow_incomplete:          # date absente sur une ligne : celle de la campagne
            camp.start, camp.end = camp.start or start, camp.end or end
            start, end = start or camp.start, end or camp.end
        if (start, end) != (camp.start, camp.end):
            fail(idx, f"dates incohérentes pour la campagne '{campaign_name}' "
                      f"({start} -> {end} au lieu de {camp.start} -> {camp.end}).")
        for g, note in row_gaps:
            who = ad_name if g == "landing_page_url" else f"campagne « {campaign_name} »"
            if note:
                who += f" ({note})"
            if who not in camp.missing.setdefault(g, []):
                camp.missing[g].append(who)

        raw_budget = get("budget_eur").replace(" ", "").replace(",", ".")
        if raw_budget:
            try:
                camp.budget_eur = float(raw_budget)
            except ValueError:
                fail(idx, f"budget_eur invalide ('{get('budget_eur')}').")

        # Emplacement
        if placement_name not in camp.placements:
            site_id = get("site_id") or None
            camp.placements[placement_name] = PlacementSpec(
                name=placement_name,
                site_id=int(site_id) if site_id and site_id.isdigit() else None,
                site_name=get("site_name"),
                width=w, height=h,
                pricing_type=PRICING_TYPES.get(get("pricing_type").lower(), "PRICING_TYPE_CPM"),
                row_index=idx,
                category=get("placement_category").lower()
                if get("placement_category").lower() in PLACEMENT_CATEGORIES else "",
                objective=get("objective"), audience=get("audience"),
                targeting=get("targeting"), format_label=get("format_label"),
                package=get("package"), device=get("device"),
                budget=get("line_budget_eur"), volume=get("line_volume"),
                extra={k: get(k) for k in EXTRA_TOKENS if get(k)},
            )
        pl_spec = camp.placements[placement_name]
        pl_spec.rows.append(idx)
        if not pl_spec.final_name and get("placement_final"):
            pl_spec.final_name = get("placement_final")

        # Créative
        ctype, auto, err = resolve_creative_type(
            get("creative_type"), get("creative_category"), get("snippet"), get("creative_path"))
        if err:
            fail(idx, err)
            continue
        creative = CreativeSpec(
            name=creative_name, ctype=ctype, auto=auto,
            snippet=get("snippet"), path=get("creative_path"),
            landing_page_url=lp or (camp.ads.get(ad_name) or {}).get("landing_page_url", ""),
            active=_as_bool(get("active") or "1"),
            row_index=idx,
            message=get("message"),
            rows=[idx], final_name=get("creative_final"),
        )
        if ctype == "REDIRECT" and not creative.snippet:
            fail(idx, "snippet manquant pour une créative REDIRECT.")
            continue
        if ctype in ("IMAGE", "VIDEO", "HTML5", "AUDIO") and not creative.path:
            fail(idx, f"creative_path manquant pour une créative {ctype}.")
            continue
        if ctype == "VIDEO" and os.path.splitext(creative.path)[1].lower() not in VIDEO_EXT:
            fail(idx, f"fichier '{creative.path}' : extension vidéo attendue pour une créative VIDEO.")
            continue
        if ctype == "IMAGE" and os.path.splitext(creative.path)[1].lower() not in IMAGE_EXT:
            fail(idx, f"fichier '{creative.path}' : extension image attendue pour une créative IMAGE.")
            continue
        if ctype == "HTML5" and os.path.splitext(creative.path)[1].lower() not in HTML5_EXT:
            fail(idx, f"fichier '{creative.path}' : archive .zip ou .adz attendue pour une créative HTML5.")
            continue
        if ctype == "AUDIO" and os.path.splitext(creative.path)[1].lower() not in AUDIO_EXT:
            fail(idx, f"fichier '{creative.path}' : fichier mp3/ogg/oga attendu pour une créative audio.")
            continue

        # Annonce (fusion des lignes partageant le même ad_name)
        ad = camp.ads.setdefault(ad_name, {"placements": [], "creatives": [], "landing_page_url": lp})
        if placement_name not in ad["placements"]:
            ad["placements"].append(placement_name)
        if (campaign_name, ad_name, creative_name, placement_name) in seen_rows:
            fail(idx, f"doublon : créative '{creative_name}' déjà déclarée pour "
                      f"l'annonce '{ad_name}' et l'emplacement '{placement_name}'.")
            continue
        seen_rows.add((campaign_name, ad_name, creative_name, placement_name))
        existing = next((c for c in ad["creatives"] if c.name == creative_name), None)
        if existing is None:
            ad["creatives"].append(creative)
        else:
            existing.rows.append(idx)
            existing.final_name = existing.final_name or creative.final_name
        if lp:
            ad["landing_page_url"] = lp

    # Catégorie d'emplacement non fournie : 1x1 -> tracking, vidéo in-stream si une
    # créative vidéo y est diffusée, sinon display (formats IAB avec leur dimension)
    for camp in plan.values():
        video = {pn for ad in camp.ads.values() if any(c.ctype == "VIDEO" for c in ad["creatives"])
                 for pn in ad["placements"]}
        audio = {pn for ad in camp.ads.values() if any(c.ctype == "AUDIO" for c in ad["creatives"])
                 for pn in ad["placements"]}
        for p in camp.placements.values():
            if not p.category:
                p.category = ("tracking" if (p.width, p.height) == (1, 1)
                              else "video" if p.name in video
                              else "audio" if p.name in audio else "display")

    # Une annonce de tracking ne peut pas mélanger tracking et créatives affichées
    for camp in plan.values():
        for ad_name, ad in camp.ads.items():
            kinds = {c.ctype == "TRACKING" for c in ad["creatives"]}
            if len(kinds) > 1:
                errors.append(f"Annonce '{ad_name}' : mélange de créatives TRACKING et "
                              "affichées/vidéo — séparez-les en deux annonces.")
    return plan, grouped(errors)


def plan_gaps(plan: Dict[str, CampaignSpec]) -> List[str]:
    """Champs laissés vides, à compléter avant l'envoi (une phrase par champ)."""
    out = []
    for camp in plan.values():
        for f, who in camp.missing.items():
            if f in ("campaign_start", "campaign_end") and getattr(camp, "start" if f == "campaign_start" else "end"):
                continue          # renseignée sur une autre ligne de la campagne
            out.append(f"{GAP_LABELS.get(f, f)} — à compléter : {', '.join(who)}")
    return out


def print_plan_summary(plan: Dict[str, CampaignSpec], errors: List[str]) -> None:
    print(f"\n{len(plan)} campagne(s) détectée(s), {len(errors)} erreur(s).")
    for camp in plan.values():
        print(f"\nCampagne : {camp.name} ({camp.start} -> {camp.end})")
        print(f"  Annonceur : {camp.advertiser_id or camp.advertiser_name}")
        for p in camp.placements.values():
            print(f"  Placement : {p.name} — {p.width}x{p.height} — {p.pricing_type}"
                  f" — site {p.site_id or p.site_name or '?'}")
        for ad_name, ad in camp.ads.items():
            print(f"  Annonce : {ad_name} -> placements {ad['placements']} — "
                  f"{len(ad['creatives'])} créative(s) : "
                  + ", ".join(f"{c.name} ({_ctype_label(c)})" for c in ad["creatives"]))
    for e in errors:
        print("  [ERREUR] " + e)


# ---------------------------------------------------------------------------
# 2bis. Nomenclatures de nommage (couche de modification avant envoi)
# ---------------------------------------------------------------------------

def load_nomenclature(path: str) -> Dict[str, dict]:
    """Charge le fichier de nomenclatures JSON (optionnel).
    Structure attendue (voir le template fourni) :
    {
      "defaults":   regles appliquees aux annonceurs sans regle dediee,
      "advertisers": { "Nom annonceur ou ID": { regles dediees } }
    }
    Regles : "campaign" | "placement" | "ad" | "creative" (motifs) et
    "options" : separator, case (upper|lower|title), strip_accents,
    max_length, forbidden (liste de caracteres a supprimer)."""
    if not path:
        return {}
    try:
        with open(path, encoding="utf-8") as f:
            naming = json.load(f)
    except (OSError, json.JSONDecodeError) as e:
        sys.exit(f"Nomenclature illisible ({path}) : {e}")
    if not isinstance(naming, dict):
        sys.exit("Nomenclature invalide : un objet JSON est attendu.")
    return naming


def _nomenclature_rules(naming: dict, camp: CampaignSpec) -> dict:
    """Renvoie les regles de l'annonceur (par nom ou par ID), sinon les
    regles par defaut."""
    advs = naming.get("advertisers", {}) or {}
    rules = (advs.get(camp.advertiser_name or "")
             or advs.get(str(camp.advertiser_id or ""))
             or naming.get("defaults", {})
             or {})
    return rules if isinstance(rules, dict) else {}


def _nomenclature_options(rules: dict, naming: dict) -> dict:
    return rules.get("options") or naming.get("defaults", {}).get("options") or {}


def _apply_nomenclature_options(value: str, opts: dict) -> str:
    """Normalise une valeur selon les options :
    'Campagne Été 2026' -> 'CAMPAGNE_ETE_2026' (separator + case upper)."""
    sep = str(opts.get("separator", "_"))
    if opts.get("strip_accents", True):
        value = unicodedata.normalize("NFKD", value)
        value = "".join(c for c in value if not unicodedata.combining(c))
    for forbidden in opts.get("forbidden", []):
        value = value.replace(str(forbidden), "")
    case = opts.get("case", "")
    if case == "upper":
        value = value.upper()
    elif case == "lower":
        value = value.lower()
    elif case == "title":
        value = value.title()
    value = re.sub(r"[ \t]+", sep, value.strip())
    value = re.sub(re.escape(sep) + r"+", sep, value)
    if opts.get("max_length"):
        value = value[: int(opts["max_length"])].rstrip(sep)
    return value


# Jetons supplémentaires des nomenclatures d'agence (trafficking sheet) : lus tels quels dans le plan
EXTRA_TOKENS = ("placement_type", "buy", "partner", "targeting_type", "targeting_detail", "asset_type",
                "tracking_type", "device", "info", "description")


def _pattern_context(camp: CampaignSpec, p: Optional[PlacementSpec] = None,
                     ad_name: str = "",
                     c: Optional[CreativeSpec] = None) -> dict:
    """Jetons disponibles dans les motifs de nommage."""
    ctx = {
        "advertiser_name": camp.advertiser_name,
        "advertiser_id": camp.advertiser_id or "",
        "campaign_name": camp.name,
        "campaign_start": camp.start,
        "campaign_end": camp.end,
        "year": (camp.start or "")[:4],
        "start_yyyymm": (camp.start or "").replace("-", ""),
    }
    if p is not None:
        ctx.update({
            "placement_name": p.name,
            "site_name": p.site_name,
            "site_id": p.site_id or "",
            "size": f"{p.width}x{p.height}",
            "pricing_type": p.pricing_type,
            "regie": p.site_name or (str(p.site_id) if p.site_id else ""),
            "objective": p.objective,
            "audience": p.audience,
            "targeting": p.targeting,
            "format": p.format_label or f"{p.width}x{p.height}",
            "format_label": p.format_label or f"{p.width}x{p.height}",
            "device": p.device,
        })
        ctx.update({k: v for k, v in p.extra.items() if k not in ctx or not ctx[k]})
    if ad_name:
        ctx["ad_name"] = ad_name
    if c is not None:
        ctx.update({"creative_name": c.name, "creative_type": c.ctype,
                    "message": c.message or c.name})
    return ctx


def _apply_pattern(pattern: str, ctx: dict, opts: dict) -> str:
    value = str(pattern)
    for key, val in ctx.items():
        value = value.replace("{" + key + "}", str(val))
    value = re.sub(r"\{[a-zA-Z0-9_]+\}", "", value)   # jetons inconnus
    # Segments vides (information absente du plan) : pas de « -- » ni de tiret final
    value = re.sub(r"\s*([-_|/])(?:\s*\1)+\s*", r"\1", value).strip(" -_|/")
    return _apply_nomenclature_options(value, opts)


# Nomenclature de base, appliquée quand aucun fichier de nomenclature n'est fourni :
#   emplacement : OBJECTIF-AUDIENCE-RÉGIE-CIBLAGE-FORMAT
#   créative    : RÉGIE-FORMAT-MESSAGE   (annonce = nom de l'emplacement)
BASE_NOMENCLATURE = {
    "defaults": {
        "placement": "{objective}-{audience}-{regie}-{targeting}-{format}",
        "ad": "{placement_name}",
        # Package CM360 = régie + objectif + ciblage (créé si le plan n'en donne pas) :
        # un package par ciblage quand une ligne du plan en décrit plusieurs
        "package": "Package-{regie}-{objective}-{targeting}",
        "creative": "{regie}-{format}-{message}",
        "options": {"separator": "_", "case": "upper", "strip_accents": True,
                    "max_length": 255, "forbidden": []},
        # Sans objectif, audience ni ciblage, les noms d'origine sont conservés
        "require_any": ["objective", "audience", "targeting"],
    }
}


def naming_or_base(naming: dict) -> dict:
    """Nomenclature fournie, sinon la nomenclature de base."""
    return naming if naming else BASE_NOMENCLATURE


def _unique(name: str, used: set) -> str:
    """Évite deux objets du même nom (ils seraient fusionnés à l'envoi)."""
    if name not in used:
        used.add(name)
        return name
    n = 2
    while f"{name}-{n}" in used:
        n += 1
    used.add(f"{name}-{n}")
    return f"{name}-{n}"


def _apply_final_names(camp: "CampaignSpec", renames: List[str]) -> None:
    """Noms imposés à la main (aperçu / tableau) sans autre nomenclature."""
    used: set = set()
    pmap: Dict[str, str] = {}
    new_pl: Dict[str, PlacementSpec] = {}
    for p in camp.placements.values():
        name = _unique(p.final_name or p.name, used)
        if name != p.name:
            renames.append(f"Emplacement : « {p.name} » -> « {name} »")
        pmap[p.name] = name
        p.name = name
        new_pl[name] = p
    camp.placements = new_pl
    for ad in camp.ads.values():
        ad["placements"] = [pmap.get(pn, pn) for pn in ad["placements"]]
        for c in ad["creatives"]:
            if c.final_name and c.final_name != c.name:
                renames.append(f"Créative : « {c.name} » -> « {c.final_name} »")
                c.name = c.final_name


def apply_nomenclature(plan: Dict[str, CampaignSpec], naming: dict,
                       log: Callable[[str], None] = print
                       ) -> Tuple[Dict[str, CampaignSpec], List[str]]:
    """Applique les nomenclatures de nommage par annonceur au plan :
    renomme campagnes, emplacements, annonces et creatives (les references
    placements des annonces sont remappees automatiquement).
    Renvoie (plan modifie, liste des renommages affiches avant l'envoi)."""
    if not naming:
        return plan, []
    renames: List[str] = []
    for camp in plan.values():
        rules = _nomenclature_rules(naming, camp)
        if not rules:
            continue
        opts = _nomenclature_options(rules, naming)

        # Campagne
        if rules.get("campaign"):
            new_name = _apply_pattern(rules["campaign"],
                                      _pattern_context(camp), opts)
            if new_name and new_name != camp.name:
                renames.append(f"Campagne : « {camp.name} » -> « {new_name} »")
                camp.name = new_name

        # Emplacements (references remappees dans les annonces ensuite)
        placement_map: Dict[str, str] = {}
        need = rules.get("require_any") or []
        if need and not any(getattr(p, f, "") for p in camp.placements.values() for f in need):
            _apply_final_names(camp, renames)   # noms d'origine, sauf noms imposés à la main
            continue
        if rules.get("package"):
            for p in camp.placements.values():
                if not p.package and p.objective:
                    p.package = _apply_pattern(rules["package"], _pattern_context(camp, p=p), opts)
        if rules.get("placement"):
            new_placements: Dict[str, PlacementSpec] = {}
            used: set = set()
            for p in camp.placements.values():
                new_name = p.final_name or _apply_pattern(rules["placement"],
                                                          _pattern_context(camp, p=p), opts)
                new_name = _unique(new_name or p.name, used)
                if new_name != p.name:
                    renames.append(f"Emplacement : « {p.name} » -> « {new_name} »")
                placement_map[p.name] = new_name
                p.name = new_name
                new_placements[new_name] = p
            camp.placements = new_placements

        # Annonces + creatives
        new_ads: Dict[str, dict] = {}
        used_ads: set = set()
        used_creas: set = set()
        crea_sig: Dict[str, tuple] = {}         # nom -> fichier / tag : même asset = même créative partagée
        renamed_creas: Dict[int, str] = {}
        for ad_name, ad in camp.ads.items():
            if placement_map:
                ad["placements"] = [placement_map.get(pn, pn)
                                     for pn in ad["placements"]]
            p0 = camp.placements.get(ad["placements"][0]) if ad["placements"] else None
            if rules.get("ad"):
                new_ad_name = _apply_pattern(
                    rules["ad"], _pattern_context(camp, p=p0, ad_name=ad_name), opts)
                new_ad_name = _unique(new_ad_name or ad_name, used_ads)
                if new_ad_name != ad_name:
                    renames.append(f"Annonce : « {ad_name} » -> « {new_ad_name} »")
                    ad_name = new_ad_name
            if rules.get("creative"):
                for c in ad["creatives"]:
                    if id(c) in renamed_creas:
                        continue          # même créative partagée par plusieurs annonces
                    ctx = _pattern_context(camp, p=p0, ad_name=ad_name, c=c)
                    wanted = c.final_name or _apply_pattern(rules["creative"], ctx, opts) or c.name
                    sig = (c.ctype, c.path, c.snippet)
                    if wanted in crea_sig and crea_sig[wanted] == sig and not c.path and not c.snippet:
                        new_c_name = wanted          # même message, sans fichier propre : une seule créative partagée
                    else:
                        new_c_name = _unique(wanted, used_creas)
                        crea_sig.setdefault(new_c_name, sig)
                    renamed_creas[id(c)] = new_c_name
                    if new_c_name and new_c_name != c.name:
                        renames.append(f"Créative : « {c.name} » -> « {new_c_name} »")
                        c.name = new_c_name
            new_ads[ad_name] = ad
        camp.ads = new_ads
    if renames and log:
        log(f"[NOMENCLATURE] {len(renames)} objet(s) renommé(s) selon les "
            "conventions de nommage de l'annonceur.")
    return plan, renames


# ---------------------------------------------------------------------------
# 3. Interface d'adaptateur (contrat commun à toutes les plateformes)
# ---------------------------------------------------------------------------

class TraffickingAdapter:
    """Contrat que chaque plateforme doit implémenter.

    L'orchestrateur (traffic_plan) appelle toujours les mêmes méthodes,
    dans le même ordre. Un adaptateur est responsable :
      - de la traduction plan média -> objets de la plateforme ;
      - de l'idempotence (réutilisation des objets existants par nom) ;
      - du rapport des objets créés / réutilisés (self.report).
    """

    name = "base"

    def __init__(self, log: Callable[[str], None] = print):
        self.log = log
        self.report: Dict[str, dict] = {"created": {}, "reused": {}, "errors": []}

    # -- suivi du rapport ------------------------------------------------

    def _record(self, kind: str, name: str, obj_id, created: bool):
        bucket = self.report["created"] if created else self.report["reused"]
        bucket.setdefault(kind, {})[name] = obj_id

    # -- méthodes à implémenter par chaque plateforme ----------------------

    def check_access(self) -> None:
        """Valide les identifiants sans rien modifier (mode --check)."""
        raise NotImplementedError

    def ensure_advertiser(self, camp: CampaignSpec) -> int:
        """Retourne l'ID de l'annonceur (créé si absent)."""
        raise NotImplementedError

    def ensure_landing_page(self, advertiser_id: int, url: str) -> int:
        raise NotImplementedError

    def ensure_campaign(self, advertiser_id: int, camp: CampaignSpec,
                        default_lp_id: int) -> int:
        raise NotImplementedError

    def ensure_placement(self, campaign_id: int, p: PlacementSpec) -> int:
        raise NotImplementedError

    def ensure_creative(self, advertiser_id: int, campaign_id: int, c: CreativeSpec,
                        width: int, height: int) -> int:
        """Retourne l'ID de la créative et l'associe à la campagne si la
        plateforme l'exige (CM360 : campaignCreativeAssociations)."""
        raise NotImplementedError

    def ensure_ad(self, campaign_id: int, ad_name: str, ad: dict,
                  placement_ids: List[int], creative_ids: Dict[str, int]) -> int:
        raise NotImplementedError


# ---------------------------------------------------------------------------
# 4. Adaptateur CM360 (Campaign Manager 360, API v5)
# ---------------------------------------------------------------------------

def cm360_authenticate(service_account_path: str):
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    creds = service_account.Credentials.from_service_account_file(
        service_account_path, scopes=CM360_SCOPES)
    return build("dfareporting", "v5", credentials=creds, cache_discovery=False)


def cm360_pick_profile(service, profile_id: Optional[int] = None) -> int:
    profiles = service.userprofiles().list().execute().get("items", [])
    if not profiles:
        sys.exit("Aucun profil utilisateur CM360 accessible. Vérifiez que l'email du "
                 "compte de service a bien été ajouté comme utilisateur dans CM360.")
    if profile_id:
        for p in profiles:
            if int(p["id"]) == int(profile_id):
                return int(p["id"])
        sys.exit(f"Profil {profile_id} introuvable. Profils disponibles : "
                 + ", ".join(f"{p['id']} ({p['accountName']})" for p in profiles))
    if len(profiles) > 1:
        print("Plusieurs profils disponibles : "
              + ", ".join(f"{p['id']} ({p['accountName']})" for p in profiles)
              + " — le premier est utilisé (sinon passez --profile-id).")
    return int(profiles[0]["id"])


def cm360_api_call(fn, log: Callable[[str], None], **kwargs):
    """Exécute un appel API CM360 avec retry (backoff) sur 429 / 5xx."""
    for attempt in range(5):
        try:
            return fn(**kwargs).execute()
        except Exception as e:
            status = getattr(e, "status_code", None)
            resp = getattr(e, "resp", None)
            if resp is not None and hasattr(resp, "status"):
                status = status or resp.status
            retryable = status in (429, 500, 502, 503) or "rate" in str(e).lower()
            if retryable and attempt < 4:
                wait = 2 ** attempt
                log(f"  ... appel API rejeté (code {status}), nouvelle tentative dans {wait}s")
                time.sleep(wait)
                continue
            raise


class CM360Adapter(TraffickingAdapter):
    """Campaign Manager 360 — API v5 (dfareporting.googleapis.com)."""
    name = "cm360"

    def __init__(self, service, profile_id: int,
                 log: Callable[[str], None] = print, creatives_dir: str = ""):
        super().__init__(log=log)
        self.svc = service
        self.profile = profile_id
        self.creatives_dir = creatives_dir
        self._sizes = None
        self._site_cache: Dict[str, int] = {}

    # -- helpers -----------------------------------------------------------

    def _call(self, fn, **kwargs):
        return cm360_api_call(fn, self.log, **kwargs)

    # -- diagnostic ----------------------------------------------------------

    def check_access(self) -> None:
        profiles = self.svc.userprofiles().list().execute().get("items", [])
        if not profiles:
            sys.exit("Aucun profil CM360 accessible. L'email du compte de service "
                     "n'est probablement pas encore ajouté comme utilisateur dans "
                     "CM360 (voir le guide, §2, étape 4).")
        print("Profils CM360 accessibles :")
        for p in profiles:
            print(f"  - {p['id']} : {p.get('accountName', '?')} "
                  f"(compte {p.get('accountId', '?')})")
        pid = self.profile
        print(f"\nDiagnostic via le profil {pid} :")
        adv = self._call(self.svc.advertisers().list, profileId=pid, maxResults=10)
        print(f"Annonceurs ({len(adv.get('items', []))} premiers) :")
        for a in adv.get("items", []):
            print(f"  - advertiser_id = {a['id']} : {a['name']}")
        sites = self._call(self.svc.sites().list, profileId=pid, maxResults=10)
        print(f"Sites ({len(sites.get('items', []))} premiers) :")
        for s in sites.get("items", []):
            print(f"  - site_id = {s['id']} : {s['name']}")
        print("\nAccès OK. Réutilisez ces IDs dans le plan média "
              "(colonnes advertiser_id et site_id).")

    # -- annonceur -----------------------------------------------------------

    def ensure_advertiser(self, camp: CampaignSpec) -> int:
        if camp.advertiser_id:
            adv = self._call(self.svc.advertisers().get,
                             profileId=self.profile, id=camp.advertiser_id)
            self._record("advertisers", adv["name"], adv["id"], False)
            return adv["id"]
        if not camp.advertiser_name:
            raise ValueError("Indiquez advertiser_id ou advertiser_name dans le plan média.")
        resp = self._call(self.svc.advertisers().list,
                          profileId=self.profile, searchString=camp.advertiser_name)
        for adv in resp.get("items", []):
            if adv["name"].lower() == camp.advertiser_name.lower():
                self._record("advertisers", adv["name"], adv["id"], False)
                return adv["id"]
        adv = self._call(self.svc.advertisers().insert, profileId=self.profile,
                         body={"name": camp.advertiser_name, "status": "APPROVED"})
        self.log(f"  Annonceur créé : {adv['name']} (id {adv['id']})")
        self._record("advertisers", adv["name"], adv["id"], True)
        return adv["id"]

    # -- landing page --------------------------------------------------------

    def ensure_landing_page(self, advertiser_id: int, url: str) -> int:
        resp = self._call(self.svc.advertiserLandingPages().list,
                          profileId=self.profile, advertiserIds=[advertiser_id],
                          searchString=url)
        for lp in resp.get("items", []):
            if lp.get("url", "").rstrip("/") == url.rstrip("/"):
                self._record("landing_pages", lp["name"], lp["id"], False)
                return lp["id"]
        name = f"LP_{url.rstrip('/').split('//')[-1][:80]}"
        lp = self._call(self.svc.advertiserLandingPages().insert,
                        profileId=self.profile,
                        body={"advertiserId": advertiser_id, "name": name, "url": url})
        self.log(f"  Landing page créée : {name} (id {lp['id']})")
        self._record("landing_pages", name, lp["id"], True)
        return lp["id"]

    # -- campagne ------------------------------------------------------------

    def ensure_campaign(self, advertiser_id: int, camp: CampaignSpec,
                        default_lp_id: int) -> int:
        resp = self._call(self.svc.campaigns().list,
                          profileId=self.profile, advertiserIds=[advertiser_id],
                          archived=False, searchString=camp.name)
        for c in resp.get("items", []):
            if c["name"].lower() == camp.name.lower():
                self._record("campaigns", c["name"], c["id"], False)
                return c["id"]
        body = {
            "advertiserId": advertiser_id,
            "name": camp.name,
            "archived": False,
            "startDate": camp.start,
            "endDate": camp.end,
            "defaultLandingPageId": default_lp_id,
        }
        c = self._call(self.svc.campaigns().insert,
                       profileId=self.profile, body=body)
        self.log(f"  Campagne créée : {c['name']} (id {c['id']})")
        self._record("campaigns", c["name"], c["id"], True)
        return c["id"]

    # -- tailles & sites -------------------------------------------------------

    def ensure_size(self, width: int, height: int) -> int:
        if self._sizes is None:
            self._sizes = {}
            resp = self._call(self.svc.sizes().list, profileId=self.profile)
            for s in resp.get("items", []):
                if s.get("width") and s.get("height"):
                    self._sizes[(s["width"], s["height"])] = s["id"]
        if (width, height) in self._sizes:
            return self._sizes[(width, height)]
        s = self._call(self.svc.sizes().insert, profileId=self.profile,
                       body={"width": width, "height": height})
        self._sizes[(width, height)] = s["id"]
        self.log(f"  Taille créée : {width}x{height} (id {s['id']})")
        return s["id"]

    def resolve_site(self, p: PlacementSpec) -> int:
        if p.site_id:
            return p.site_id
        if not p.site_name:
            raise ValueError(f"Placement '{p.name}' : site_id ou site_name requis.")
        key = p.site_name.lower()
        if key in self._site_cache:
            return self._site_cache[key]
        resp = self._call(self.svc.sites().list,
                          profileId=self.profile, searchString=p.site_name)
        items = resp.get("items", [])
        match = next((s for s in items if s["name"].lower() == key), None)
        if match is None:
            raise ValueError(f"Site '{p.site_name}' introuvable (ligne {p.row_index + 2}).")
        self._site_cache[key] = match["id"]
        return match["id"]

    # -- emplacements -----------------------------------------------------------

    def ensure_placement(self, campaign_id: int, p: PlacementSpec) -> int:
        resp = self._call(self.svc.placements().list,
                          profileId=self.profile, campaignIds=[campaign_id],
                          searchString=p.name)
        for pl in resp.get("items", []):
            if pl["name"].lower() == p.name.lower():
                self._record("placements", p.name, pl["id"], False)
                return pl["id"]
        site_id = self.resolve_site(p)
        size_id = self.ensure_size(p.width, p.height)
        body = {
            "campaignId": campaign_id,
            "name": p.name,
            "siteId": site_id,
            "sizeId": size_id,
            "pricingType": p.pricing_type,
            "paymentSource": "PLACEMENT_AGENCY_PAID",
            "tagSettings": {"keywordOption":
                            "GENERATE_SEPARATE_TAG_FOR_EACH_COMPLIANT_AD_SERVER"},
        }
        pl = self._call(self.svc.placements().insert,
                        profileId=self.profile, body=body)
        self.log(f"  Placement créé : {p.name} (id {pl['id']})")
        self._record("placements", p.name, pl["id"], True)
        return pl["id"]

    # -- créatives ---------------------------------------------------------------

    def ensure_creative(self, advertiser_id: int, campaign_id: int, c: CreativeSpec,
                        width: int, height: int) -> int:
        resp = self._call(self.svc.creatives().list,
                          profileId=self.profile, campaignIds=[campaign_id],
                          searchString=c.name)
        for cr in resp.get("items", []):
            if cr["name"].lower() == c.name.lower():
                self._record("creatives", c.name, cr["id"], False)
                return cr["id"]
        if c.ctype == "REDIRECT":
            body = {
                "advertiserId": advertiser_id,
                "name": c.name,
                "type": "THIRD_PARTY_REDIRECT",
                "active": c.active,
                "width": width,
                "height": height,
                "snippet": c.snippet,
                "backupImageClickThroughUrl": c.landing_page_url,
            }
        elif c.ctype == "TRACKING":
            # Créative de tracking : aucun fichier, aucune taille. La destination
            # est portée par l'annonce (AD_SERVING_TRACKING). Types CM360 à valider
            # avec un accès réel.
            body = {
                "advertiserId": advertiser_id,
                "name": c.name,
                "type": "TRACKING_TEXT",
                "active": c.active,
            }
        elif c.ctype == "VIDEO":
            # Vidéo in-stream : fichier uploadé comme asset VIDEO (envoi reprenable).
            from googleapiclient.http import MediaFileUpload
            full_path = resolve_creative_path(c.path, self.creatives_dir)
            if not os.path.isfile(full_path):
                raise ValueError(f"Créative VIDEO '{c.name}' : fichier introuvable ({full_path}).")
            asset = self._call(
                self.svc.creativeAssets().insert,
                profileId=self.profile, advertiserId=advertiser_id,
                body={"assetIdentifier": {"name": os.path.basename(full_path),
                                          "type": "VIDEO"}},
                media_body=MediaFileUpload(full_path, resumable=True))
            body = {
                "advertiserId": advertiser_id,
                "name": c.name,
                "type": "INSTREAM_VIDEO",
                "active": c.active,
                "creativeAssets": [{
                    "assetIdentifier": asset["assetIdentifier"],
                    "role": "PRIMARY",
                }],
                "backupImageClickThroughUrl": c.landing_page_url,
            }
        elif c.ctype in ("HTML5", "AUDIO"):
            # Contrôlées selon les specs CM360 (creative_specs.py) ; création par l'API
            # non encore validée sur un vrai compte : passer par l'export Campaign
            # spreadsheet / le Bulk Editor, ou l'upload manuel.
            raise ValueError(f"Créative {c.ctype} '{c.name}' : création via l'API CM360 pas encore "
                             "disponible — utilisez l'export Campaign spreadsheet (Bulk Editor).")
        else:  # IMAGE (display)
            full_path = resolve_creative_path(c.path, self.creatives_dir)
            if not os.path.isfile(full_path):
                raise ValueError(f"Créative IMAGE '{c.name}' : fichier introuvable ({full_path}).")
            asset = self._call(
                self.svc.creativeAssets().insert,
                profileId=self.profile, advertiserId=advertiser_id,
                body={"assetIdentifier": {"name": os.path.basename(full_path),
                                          "type": "IMAGE"}},
                media_body=full_path)
            body = {
                "advertiserId": advertiser_id,
                "name": c.name,
                "type": "IMAGE",
                "active": c.active,
                "width": width,
                "height": height,
                "artworkType": "ARTWORK_TYPE_IMAGE",
                "creativeAssets": [{
                    "assetIdentifier": asset["assetIdentifier"],
                    "size": {"width": width, "height": height},
                    "role": "PRIMARY",
                }],
                "backupImageClickThroughUrl": c.landing_page_url,
            }
        cr = self._call(self.svc.creatives().insert,
                        profileId=self.profile, body=body)
        self.log(f"  Créative créée : {c.name} ({_ctype_label(c)}, id {cr['id']})")
        self._record("creatives", c.name, cr["id"], True)

        # Obligatoire : associer la créative à la campagne AVANT l'annonce.
        self._call(self.svc.campaignCreativeAssociations().insert,
                   profileId=self.profile, campaignId=campaign_id,
                   body={"creativeId": cr["id"]})
        return cr["id"]

    # -- annonces ------------------------------------------------------------------

    def ensure_ad(self, campaign_id: int, ad_name: str, ad: dict,
                 placement_ids: List[int], creative_ids: Dict[str, int]) -> int:
        resp = self._call(self.svc.ads().list,
                          profileId=self.profile, campaignIds=[campaign_id],
                          searchString=ad_name)
        for a in resp.get("items", []):
            if a["name"].lower() == ad_name.lower():
                self._record("ads", ad_name, a["id"], False)
                return a["id"]
        body = {
            "campaignId": campaign_id,
            "name": ad_name,
            "active": True,
            "type": ("AD_SERVING_TRACKING"
                     if all(c.ctype == "TRACKING" for c in ad["creatives"])
                     else "AD_SERVING_STANDARD_AD"),
            "creativeRotation": {
                "type": "CREATIVE_ROTATION_TYPE_RANDOM",
                "creativeAssignments": [
                    {"creativeId": creative_ids[c.name], "active": c.active}
                    for c in ad["creatives"]
                ],
            },
            "placementAssignments": [
                {"placementId": pid, "active": True} for pid in placement_ids
            ],
            "creativeClickThroughUrl": ad["landing_page_url"],
        }
        a = self._call(self.svc.ads().insert, profileId=self.profile, body=body)
        self.log(f"  Annonce créée : {ad_name} (id {a['id']})")
        self._record("ads", ad_name, a["id"], True)
        return a["id"]


# ---------------------------------------------------------------------------
# 5. Adaptateur Adform (api.adform.com, OAuth 2.0 client credentials)
# ---------------------------------------------------------------------------

class AdformAdapter(TraffickingAdapter):
    """Adform — API externe (https://api.adform.com).

    ÉTAT : authentification et diagnostic (--check) implémentés.
    La création d'objets (campagnes, line items, bannières) sera complétée
    à l'obtention d'un accès réel : l'API externe doit être activée sur le
    compte client Adform, et l'OAuth client est délivré par le support
    (technical@adform.com). Mapping prévu :
        annonceur -> advertiser    campagne -> campaign
        emplacement -> line item  créative -> banner
    """

    name = "adform"

    def __init__(self, client_id: str, client_secret: str,
                 log: Callable[[str], None] = print):
        super().__init__(log=log)
        self.client_id = client_id
        self.client_secret = client_secret
        self._token: Optional[str] = None

    # -- HTTP / OAuth --------------------------------------------------------

    def _get_token(self) -> str:
        """OAuth 2.0 client credentials contre id.adform.com (stdlib only)."""
        import urllib.parse
        import urllib.request
        if self._token:
            return self._token
        data = urllib.parse.urlencode({
            "grant_type": "client_credentials",
            "client_id": self.client_id,
            "client_secret": self.client_secret,
            "scope": ADFORM_SCOPE,
        }).encode()
        req = urllib.request.Request(
            ADFORM_TOKEN_URL, data=data,
            headers={"Content-Type": "application/x-www-form-urlencoded"})
        try:
            with urllib.request.urlopen(req) as resp:
                payload = json.loads(resp.read().decode())
        except Exception as e:
            raise RuntimeError(f"Authentification Adform échouée : {e}") from e
        self._token = payload["access_token"]
        return self._token

    def _request(self, method: str, path: str, data: Optional[dict] = None):
        import urllib.request
        req = urllib.request.Request(
            ADFORM_API_BASE + path,
            data=json.dumps(data).encode() if data is not None else None,
            headers={"Authorization": f"Bearer {self._get_token()}",
                     "Content-Type": "application/json"},
            method=method)
        with urllib.request.urlopen(req) as resp:
            body = resp.read().decode()
            return json.loads(body) if body else {}

    # -- diagnostic ------------------------------------------------------------

    def check_access(self) -> None:
        camps = self._request("GET", "/v1/buyer/campaigns")
        count = len(camps) if isinstance(camps, list) else "?"
        self.log(f"Accès Adform OK — {count} campagne(s) visible(s).")
        self.log("Rappel : la création d'objets via l'API Adform sera complétée "
                 "dès l'obtention d'un compte de test (voir docstring).")

    # -- création d'objets : à compléter avec un accès réel ---------------------

    def ensure_advertiser(self, camp: CampaignSpec) -> int:
        raise NotImplementedError(
            "AdformAdapter.ensure_advertiser : à implémenter avec l'accès réel "
            "(endpoint advertisers de l'API Adform).")

    def ensure_landing_page(self, advertiser_id: int, url: str) -> int:
        raise NotImplementedError(
            "AdformAdapter.ensure_landing_page : à implémenter avec l'accès réel.")

    def ensure_campaign(self, advertiser_id: int, camp: CampaignSpec,
                        default_lp_id: int) -> int:
        raise NotImplementedError(
            "AdformAdapter.ensure_campaign : à implémenter avec l'accès réel "
            "(POST /v1/buyer/campaigns).")

    def ensure_placement(self, campaign_id: int, p: PlacementSpec) -> int:
        raise NotImplementedError(
            "AdformAdapter.ensure_placement : à implémenter avec l'accès réel "
            "(line items de l'API Adform).")

    def ensure_creative(self, advertiser_id: int, campaign_id: int, c: CreativeSpec,
                        width: int, height: int) -> int:
        raise NotImplementedError(
            "AdformAdapter.ensure_creative : à implémenter avec l'accès réel "
            "(banners de l'API Adform). Typage déjà fourni par le plan : "
            f"{c.ctype} (display / vidéo in-stream / tracking).")

    def ensure_ad(self, campaign_id: int, ad_name: str, ad: dict,
                  placement_ids: List[int], creative_ids: Dict[str, int]) -> int:
        raise NotImplementedError(
            "AdformAdapter.ensure_ad : à implémenter avec l'accès réel.")


# ---------------------------------------------------------------------------
# 6. Adaptateur démo (mock) : simulation locale, aucun accès requis
# ---------------------------------------------------------------------------

class MockAdapter(TraffickingAdapter):
    """Simule n'importe quelle plateforme : mêmes appels, IDs fictifs.
    Sert au développement, aux démos et aux tests de bout en bout."""
    name = "mock"

    def __init__(self, log: Callable[[str], None] = print, creatives_dir: str = ""):
        super().__init__(log=log)
        self.creatives_dir = creatives_dir
        self._ids: Dict[Tuple[str, str], int] = {}
        self._next = 5000

    def _mock_id(self, kind: str, name: str, created_msg: str) -> int:
        key = (kind, name)
        if key not in self._ids:
            self._next += 1
            self._ids[key] = self._next
            self.log(f"  {created_msg} (id {self._next}) [simulé]")
            self._record(kind, name, self._next, True)
        else:
            self._record(kind, name, self._ids[key], False)
        return self._ids[key]

    def check_access(self) -> None:
        self.log("Profil MOCK (démo) : accès simulé, aucun compte requis.")
        self.log("Campagnes fictives : aucune. Annonceurs fictifs : aucun.")

    def ensure_advertiser(self, camp: CampaignSpec) -> int:
        name = camp.advertiser_name or f"Annonceur {camp.advertiser_id}"
        return self._mock_id("advertisers", name, f"Annonceur créé : {name}")

    def ensure_landing_page(self, advertiser_id: int, url: str) -> int:
        return self._mock_id("landing_pages", url, f"Landing page créée : {url}")

    def ensure_campaign(self, advertiser_id: int, camp: CampaignSpec,
                        default_lp_id: int) -> int:
        return self._mock_id("campaigns", camp.name,
                              f"Campagne créée : {camp.name} "
                              f"({camp.start} -> {camp.end})")

    def ensure_placement(self, campaign_id: int, p: PlacementSpec) -> int:
        return self._mock_id("placements", p.name,
                              f"Placement créé : {p.name} ({p.width}x{p.height})")

    def ensure_creative(self, advertiser_id: int, campaign_id: int, c: CreativeSpec,
                        width: int, height: int) -> int:
        # Mock : le fichier image n'est pas lu (simulation), pas de contrôle d'existence.
        return self._mock_id("creatives", c.name,
                             f"Créative créée : {c.name} ({_ctype_label(c)}, {width}x{height})")

    def ensure_ad(self, campaign_id: int, ad_name: str, ad: dict,
                  placement_ids: List[int], creative_ids: Dict[str, int]) -> int:
        return self._mock_id("ads", ad_name,
                              f"Annonce créée : {ad_name} "
                              f"({len(creative_ids)} créative(s), "
                              f"{len(placement_ids)} placement(s))")


# ---------------------------------------------------------------------------
# 7. Fabrique d'adaptateurs
# ---------------------------------------------------------------------------

def build_adapter(platform: str = "cm360",
                  log: Callable[[str], None] = print,
                  service_account: str = "",
                  profile_id: Optional[int] = None,
                  adform_client_id: str = "",
                  adform_client_secret: str = "",
                  creatives_dir: str = "") -> TraffickingAdapter:
    """Construit l'adaptateur correspondant à la plateforme demandée.
    Choix : cm360 | adform | mock."""
    if platform in ("mock", "demo"):
        return MockAdapter(log=log, creatives_dir=creatives_dir)
    if platform == "cm360":
        if not service_account:
            sys.exit("CM360 : compte de service requis (--service-account cle.json "
                     "ou CM360_SA_KEY). Sans accès, testez avec --platform mock.")
        service = cm360_authenticate(service_account)
        pid = cm360_pick_profile(service, profile_id)
        return CM360Adapter(service, pid, log=log, creatives_dir=creatives_dir)
    if platform == "adform":
        if not (adform_client_id and adform_client_secret):
            sys.exit("Adform : identifiants OAuth requis (--adform-client-id / "
                     "--adform-client-secret, ou ADFORM_CLIENT_ID / "
                     "ADFORM_CLIENT_SECRET).")
        return AdformAdapter(adform_client_id, adform_client_secret, log=log)
    sys.exit(f"Plateforme inconnue : '{platform}' (cm360 | adform | mock)")


# ---------------------------------------------------------------------------
# 8. Orchestration (indépendante de la plateforme)
# ---------------------------------------------------------------------------

def traffic_plan(adapter: TraffickingAdapter, plan: Dict[str, CampaignSpec],
                 log: Callable[[str], None] = print) -> Dict[str, dict]:
    """Exécute le trafficking complet via l'adaptateur. Renvoie le rapport."""
    for camp in plan.values():
        log(f"\n=== Campagne : {camp.name} [plateforme : {adapter.name}] ===")
        gaps = plan_gaps({camp.name: camp})
        if gaps:
            if adapter.name != "mock":      # jamais d'envoi réel d'un plan incomplet
                msg = (f"Campagne '{camp.name}' non envoyée : informations à compléter "
                       f"({'; '.join(gaps)}).")
                log("[ERREUR] " + msg)
                adapter.report["errors"].append(msg)
                continue
            for g in gaps:
                log("  [À COMPLÉTER] " + g)
        try:
            advertiser_id = adapter.ensure_advertiser(camp)
            default_lp = adapter.ensure_landing_page(advertiser_id, next(
                (ad["landing_page_url"] for ad in camp.ads.values()
                 if ad.get("landing_page_url")), "https://www.example.com"))
            campaign_id = adapter.ensure_campaign(advertiser_id, camp, default_lp)

            placement_ids = {p.name: adapter.ensure_placement(campaign_id, p)
                             for p in camp.placements.values()}

            for ad_name, ad in camp.ads.items():
                width = next((camp.placements[p].width for p in ad["placements"]), 0)
                height = next((camp.placements[p].height for p in ad["placements"]), 0)
                creative_ids = {}
                for c in ad["creatives"]:
                    creative_ids[c.name] = adapter.ensure_creative(
                        advertiser_id, campaign_id, c, width, height)
                adapter.ensure_ad(campaign_id, ad_name, ad,
                                  [placement_ids[p] for p in ad["placements"]],
                                  creative_ids)
        except Exception as e:
            msg = f"Campagne '{camp.name}' en échec : {e}"
            log("[ERREUR] " + msg)
            adapter.report["errors"].append(msg)
    return adapter.report


# ---------------------------------------------------------------------------
# 9. Point d'entrée CLI
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(
        description="Génère le trafficking (CM360 / Adform / démo) "
                    "depuis un plan média (CSV/XLSX).")
    ap.add_argument("plan", nargs="?", default="",
                    help="Fichier du plan média (.csv ou .xlsx)")
    ap.add_argument("--platform", choices=["cm360", "adform", "mock"],
                    default="cm360", help="Plateforme cible (défaut : cm360)")
    ap.add_argument("--mock", action="store_true",
                    help="Alias de --platform mock (mode démo simulé)")
    ap.add_argument("--service-account", default=os.environ.get("CM360_SA_KEY", ""),
                    help="Chemin du JSON du compte de service Google (CM360)")
    ap.add_argument("--profile-id", type=int,
                    default=int(os.environ["CM360_PROFILE_ID"])
                    if os.environ.get("CM360_PROFILE_ID") else None,
                    help="ID du profil utilisateur CM360 (défaut : 1er profil)")
    ap.add_argument("--adform-client-id",
                    default=os.environ.get("ADFORM_CLIENT_ID", ""),
                    help="Client ID OAuth Adform")
    ap.add_argument("--adform-client-secret",
                    default=os.environ.get("ADFORM_CLIENT_SECRET", ""),
                    help="Client secret OAuth Adform")
    ap.add_argument("--creatives-dir", default="",
                    help="Dossier de base des créatives IMAGE (les chemins "
                         "creative_path du plan y sont relatifs)")
    ap.add_argument("--nomenclature",
                    default=os.environ.get("CM360_NOMENCLATURE", ""),
                    help="Fichier JSON de nomenclatures de nommage par "
                         "annonceur, appliquées à la vérification avant envoi")
    ap.add_argument("--dry-run", action="store_true",
                    help="Valide et affiche le plan sans appeler l'API")
    ap.add_argument("--check", action="store_true",
                    help="Diagnostic de l'accès à la plateforme (aucune écriture)")
    ap.add_argument("--campaign-name", default="",
                    help="Nom de campagne (import d'un Campaign spreadsheet CM360, "
                         "qui ne le contient pas)")
    ap.add_argument("--advertiser-name", default="",
                    help="Nom d'annonceur (import d'un Campaign spreadsheet CM360)")
    ap.add_argument("--export-bulk", default="", metavar="FICHIER.xlsx",
                    help="Exporte le plan validé au format Campaign spreadsheet CM360 "
                         "(Bulk Editor), sans appel API")
    ap.add_argument("--export-adform-bulk", default="", metavar="FICHIER.xlsx",
                    help="Exporte campagnes et commandes au format bulk Adform (classeur v6), "
                         "sans appel API")
    ap.add_argument("--adform-campaign-type", default="",
                    help="Type de campagne Adform (défaut : Non-programmatic)")
    ap.add_argument("--report", default="",
                    help="Fichier JSON où écrire le rapport d'IDs créés")
    args = ap.parse_args()

    if args.mock:
        args.platform = "mock"

    if args.check:
        adapter = build_adapter(
            platform=args.platform, log=print,
            service_account=args.service_account, profile_id=args.profile_id,
            adform_client_id=args.adform_client_id,
            adform_client_secret=args.adform_client_secret,
            creatives_dir=args.creatives_dir)
        adapter.check_access()
        return

    if not args.plan:
        sys.exit("Indiquez le fichier du plan média (.csv ou .xlsx), "
                 "ou utilisez --check pour tester l'accès.")

    src = resolve_plan_source(args.plan)
    try:
        rows = load_rows_from_path(src, args.campaign_name, args.advertiser_name)
    finally:
        # RGPD : purge immédiate du fichier temporaire (export Google Sheet)
        if src != args.plan:
            os.remove(src)
    print(f"{len(rows)} ligne(s) lues dans {args.plan}")
    for line in getattr(rows, "ingest_report", []):
        print("  [INGESTION] " + line)
    # Préparation (dry-run, exports, démo) : les informations absentes restent vides,
    # listées « à compléter ». Envoi réel : elles bloquent (voir traffic_plan).
    allow_incomplete = bool(args.dry_run or args.export_bulk or args.export_adform_bulk
                            or args.platform == "mock")
    plan, errors = plan_from_rows(rows, allow_incomplete=allow_incomplete)
    for g in plan_gaps(plan):
        print("  [À COMPLÉTER] " + g)

    # Couche de modification à la vérification (avant envoi vers l'ad
    # server) : nomenclatures de nommage propres à chaque annonceur.
    renames: List[str] = []
    if not errors:
        naming = naming_or_base(load_nomenclature(args.nomenclature))
        if naming:
            plan, renames = apply_nomenclature(plan, naming)
            for r in renames:
                print("  [NOMENCLATURE] " + r)

    print_plan_summary(plan, errors)

    # Contrôle des fichiers créatifs selon la base de specs CM360
    spec_errors = 0
    if not errors:
        from creative_specs import check_plan
        for f in check_plan(plan, args.creatives_dir):
            tag = {"erreur": "ERREUR", "alerte": "ALERTE", "ok": "OK"}[f["level"]]
            print(f"  [SPECS CM360 {tag}] {f['message']}")
            spec_errors += f["level"] == "erreur"

    if errors:
        sys.exit("\nCorrigez les erreurs ci-dessus avant de lancer le trafficking.")
    if spec_errors and not args.dry_run and not args.export_bulk and not args.export_adform_bulk:
        sys.exit(f"\n{spec_errors} créative(s) hors specs CM360 : corrigez-les avant l'envoi.")

    if args.export_bulk:
        from cm360_bulk import write_bulk_xlsx
        n = write_bulk_xlsx(plan, args.export_bulk)
        print(f"\nCampaign spreadsheet CM360 écrit : {args.export_bulk} ({n} ligne(s)). "
              "Aucun appel API effectué.")
        return

    if args.export_adform_bulk:
        from adform_bulk import write_adform_bulk
        n = write_adform_bulk(plan, args.export_adform_bulk, args.adform_campaign_type)
        print(f"\nClasseur bulk Adform écrit : {args.export_adform_bulk} ({n} campagne(s) + commandes). "
              "Placements/annonces/créatives non exportés (voir adform_bulk.py). Aucun appel API.")
        return

    if args.dry_run:
        print(f"\n[dry-run] Aucun appel API effectué. Plan valide "
              f"(plateforme cible : {args.platform}).")
        return

    if args.platform == "mock":
        print("[MODE DÉMO] Plateforme simulée — aucun accès requis.\n")

    adapter = build_adapter(
        platform=args.platform, log=print,
        service_account=args.service_account, profile_id=args.profile_id,
        adform_client_id=args.adform_client_id,
        adform_client_secret=args.adform_client_secret,
        creatives_dir=args.creatives_dir)

    report = traffic_plan(adapter, plan)
    print(f"\nTerminé [{adapter.name}]. {len(report['errors'])} erreur(s). "
          f"Objets créés : {sum(len(v) for v in report['created'].values())}, "
          f"réutilisés : {sum(len(v) for v in report['reused'].values())}")
    if args.report:
        with open(args.report, "w", encoding="utf-8") as f:
            json.dump(report, f, ensure_ascii=False, indent=2)
        print(f"Rapport écrit : {args.report}")


if __name__ == "__main__":
    main()

