"""Trafficking sheet d'agence fictive (même structure qu'un PM d'agence réel) : onglet modèle
avec les noms construits par formule, PM avec un long bloc budget au-dessus du tableau
(en-tête ligne 45), onglet de listes (vocabulaire) et casse écrite ailleurs."""
from openpyxl import Workbook
from openpyxl.styles import PatternFill

YELLOW = PatternFill("solid", fgColor="FFFFFF00")


def build(path):
    wb = Workbook()
    m = wb.active
    m.title = "Package & Placement"
    labels = ["Type Emplacement", "Objectif Emplacement", "Mode d'Achat", "Régie", "Audience", "Ciblage",
              "Ciblage Complémentaire (non obligatoire)", "Nom Format", "Taille Format", "Device", "Description Emplacement"]
    for i, l in enumerate(labels):
        m.cell(13, 4 + i, l).fill = YELLOW
    m["R13"] = "Placement Name"
    m["R14"] = '=D14&"-"&E14&"-"&F14&"-"&G14&"-"&H14&"-"&I14&"-"&J14&"-"&K14&"-"&L14&"-"&M14&"-"&N14'
    for i, l in enumerate(["Type Asset", "Nom Format", "Taille Format", "Nom Message"]):
        m.cell(34, 4 + i, l)
    m["R34"] = "Creative Name"
    m["R35"] = '=D35&"-"&E35&"-"&F35&"-"&G35'

    pm = wb.create_sheet("PM")
    pm["B2"], pm["D2"] = "Annonceur :", "ANNONCEUR TEST"
    pm["B3"], pm["D3"] = "Campagne :", "Campagne Test"
    for r in range(13, 40):                        # bloc budget (UO, frais…) au-dessus du tableau
        pm.cell(r, 2, f"UO {r} frais"), pm.cell(r, 5, r * 10)
    head = ["Objectif", "Audience", "Support / DSP", "Régie", "Formats", "Nb \nmessages", "Détails formats",
            "Emplacements", "Devices", "Ciblages", "Date Début", "Date Fin", "Mode \nd'achat", "Volume",
            "Budget Net", "Type adserving"]
    for i, h in enumerate(head):
        pm.cell(45, 2 + i, h).fill = YELLOW
    lines = [
        ["Notoriété", "Grand public", "Amazon DSP", "AmplifiProgrammatic", "PreRoll", "POV JEUNES A\nPOV JEUNES B",
         "Preroll_20s", "RG ciblée", "Multidevice", "Netflix\nJeune\n#SocioDemo : HF1834&Geoloc",
         "06/07/2026", "26/07/2026", "CPM", 1000000, 30000, "Classique"],
        ["Notoriété", "Senior", "TF1 +", "TF1 +", "PreRoll", "CECI RETRAITES A", "Preroll_20s", "RG ciblée",
         "Multidevice", "Retraité\n#Data : DataFoyerRetraite&Géoloc", "06/07/2026", "26/07/2026", "CPM",
         700000, 15000, "Video"],
        ["Notoriété", "Grand public", "Meta (Application taxe)", "Facebook & Instagram", "LinkAdVideo",
         "POV JEUNES A\nPOV JEUNES B", "LinkAdVideo_20s (feed story reels)", "RG ciblée", "Multidevice",
         "Jeune\n#Sociodemo : HF63+\n#LAL : VIEWERS", "06/07/2026", "26/07/2026", "dCPM", 3000000, 9000, "Clics"],
        ["Considération", "Jeunes", "Amnet", "Amnet", "Display", "", "Pavé - GrandAngle", "RG ciblée", "Multidevice",
         "#Sociodemo : HF1825\n#Comportementale : JeuneDiplome", "06/07/2026", "26/07/2026", "dCPM", 500000, 4000, "Classique"],
        ["Notoriété", "Parent", "Amazon DSP", "AmplifiProgrammatic", "PreRoll", "cf. onglet", "TF1\nVideo_30s", "RG ciblée",
         "Multidevice", "#Affinitaire : PackFamille", "06/07/2026", "26/07/2026", "CPM", 300000, 6000, "Video"],
    ]
    for r, line in enumerate(lines, start=46):
        for i, v in enumerate(line):
            pm.cell(r, 2 + i, v).fill = YELLOW

    v = wb.create_sheet("Feuil2")
    cols = {"Type Emplacement": ["R", "G", "F"], "Objectif Emplacement": ["VISIBILITE", "NOTORIETE", "CONSIDERATION"],
            "Mode d'Achat": ["CPM", "CPC", "dCPM"], "Régie": ["NETFLIX", "TF1", "META", "TEADSMEDIA", "AMNET", "AMAZON"],
            "Audience": ["GRANDPUBLIC", "SENIORS", "JEUNES", "PARENTS"],
            "Ciblage": ["SOCIODEMO", "LOOKALIKE", "CUSTOM", "INTERETS", "COMPORTEMENTAL", "AFFINITAIRE"],
            "Nom Format": ["PreRoll", "LinkAdVideo", "InRead", "Pave", "GrandAngle"], "Taille Format": ["0x0", "1x1", "300x250", "300x600"],
            "Device": ["Multidevice", "Mobile"], "Type Asset": ["P", "V", "I"]}
    for c, (label, vals) in enumerate(cols.items(), start=1):
        v.cell(4, c, label)
        for r, x in enumerate(vals, start=5):
            v.cell(r, c, x)
    cache = wb.create_sheet("_cache")                # casse écrite (comme une feuille de cache Google)
    for c, x in enumerate(["Notoriete", "Consideration", "GrandPublic", "Seniors", "SocioDemo", "LookAlike", "Custom", "Comportemental"], start=1):
        cache.cell(1, c, x)
    wb.save(path)
