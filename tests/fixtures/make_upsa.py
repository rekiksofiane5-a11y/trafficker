"""Génère le plan UPSA (exemple réel anonymisé partagé pour le premier test) : cases jaunes, cellules fusionnées."""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import PatternFill

YELLOW = PatternFill("solid", fgColor="FFFFFF00")
TARGETING = ("Socio-demo : H/F – 18-49 ans \n\n"
             "Ciblage contextuel: (ciblage des mots clés, champs sémantiques, thématiques présents dans les contenus "
             "éditoriaux) : Ballonnements repas, lourdeurs estomac, maux estomac, crise de foie, nausées alcool, "
             "digestion difficile, affections digestives, remède gueule de bois, soigner gueule de bois, repas trop "
             "arrosé, repas de fête, recette pâques, repas pâques, repas festif, soirée trop arrosée, médicament "
             "digestion, aide à la digestion …\n\n"
             "+ ciblage sur les contenus mentionnant les concurrents : Charbon de Belloc, Oxyboldine, Carbolevure, "
             "Ultralevure, Gaviscon, Smecta, Spasfon, Immodium …\n"
             "Ciblage comportemental : Identifiés ayant des problèmes de digestion & nausée / intérêt contenus "
             "Santé & bien-être / Intérêt Cuisine, Repas de fête")
FORMAT = ("Bannières IAB\n(300*250)\n(160*600)\n(728*90) \n Idéalement prévoir un (300*600) pour maximiser "
          "la visibilité ")


def build(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan"
    for r, (k, v) in enumerate([("Marque", "UPSA"), ("Produit", "Citrate de Bétaine"),
                                ("Nom de campagne", "Citrate programmatique 2025"),
                                ("KPI de campagne", "Achat au CPM")], start=1):
        ws.cell(r, 1, k).fill = YELLOW
        ws.cell(r, 2, v).fill = YELLOW
    head = ["Régie", "Objectif", "KPI", "Emplacement", "Format", "Messages", "Device", " Dates de campagnes ",
            None, "Ciblage", "Mode d'achat"]
    for c, h in enumerate(head, start=1):
        ws.cell(9, c, h).fill = YELLOW
        ws.cell(10, c).fill = YELLOW
    ws.merge_cells("H9:I9")
    ws.merge_cells("A9:A10")
    row = [None, "Reach & Visibilité", "Imp", "Rotation Générale  (RON)", FORMAT, "1 message", "Multidevices",
           "16 au 23 avril 2025", None, TARGETING, "dCPM (impressions)"]
    for c, v in enumerate(row, start=1):
        cell = ws.cell(10, c, v)
        cell.fill = YELLOW
    ws.merge_cells("H10:I10")
    wb.save(path)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(__file__), "plan_upsa.xlsx"))
