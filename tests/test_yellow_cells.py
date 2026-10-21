"""Mode cases jaunes : l'agent ne lit que les cases surlignées en jaune."""
import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.styles import PatternFill

import plan_agent
from cm360_trafficker import load_rows_from_path, plan_from_rows

YELLOW = PatternFill("solid", fgColor="FFFF00")
LIGHT_YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREEN = PatternFill("solid", fgColor="C6EFCE")


def _book(tmp_path, rows, fills):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plan"
    for r in rows:
        ws.append(r)
    for (r, c), fill in fills.items():
        ws.cell(row=r, column=c).fill = fill
    p = tmp_path / "jaune.xlsx"
    wb.save(p)
    return str(p)


ROWS = [
    ["Annonceur :", "Client Démo"],
    ["Campagne", "Hiver 2026"],
    [],
    ["Support", "Emplacement", "Format", "Début", "Fin", "URL", "Commentaire"],
    ["Le Monde", "Home", "300x250", "2026-11-02", "2026-12-15", "https://x.fr", "https://autre.fr"],
    ["Le Monde", "Article", "728x90", "2026-11-02", "2026-12-15", "https://x.fr", ""],
    ["L'Équipe", "Preroll", "640x360", "2026-11-02", "2026-12-15", "https://x.fr", ""],
    ["Marmiton", "Recette", "300x600", "2026-11-02", "2026-12-15", "https://x.fr", ""],
]


def _yellow_rows(rows, cols=range(1, 7), fill=YELLOW):
    return {(r, c): fill for r in rows for c in cols}


def test_only_yellow_rows_are_read(tmp_path):
    fills = {(1, 2): YELLOW, (2, 2): YELLOW}                     # valeurs du bloc d'infos
    fills.update(_yellow_rows([5, 7]))                           # 2 lignes sur 4
    rows = load_rows_from_path(_book(tmp_path, ROWS, fills))
    plan, errors = plan_from_rows(rows)
    assert errors == []
    camp = next(iter(plan.values()))
    assert (camp.name, camp.advertiser_name) == ("Hiver 2026", "Client Démo")
    assert set(camp.placements) == {"Home", "Preroll"}
    report = "\n".join(rows.ingest_report)
    assert "Mode cases jaunes" in report


def test_non_yellow_cell_in_a_yellow_row_is_ignored(tmp_path):
    fills = {(1, 2): YELLOW, (2, 2): YELLOW}
    fills.update(_yellow_rows([5], cols=range(1, 6)))            # URL (col 6) non surlignée
    rows = load_rows_from_path(_book(tmp_path, ROWS, fills))
    assert not rows[0].get("landing_page_url")                   # ni col URL ni commentaire
    assert plan_from_rows(rows)[1]                               # l'envoi est bloqué


def test_campaign_not_highlighted_is_not_read(tmp_path):
    fills = {(1, 2): YELLOW}                                     # campagne non surlignée
    fills.update(_yellow_rows([5]))
    rows = load_rows_from_path(_book(tmp_path, ROWS, fills))
    assert rows[0]["campaign_name"] == "jaune"                  # pas « Hiver 2026 » : nom du fichier
    assert any("à vérifier" in l for l in rows.ingest_report)
    rows = load_rows_from_path(_book(tmp_path, ROWS, fills), campaign_name="Saisie")
    assert rows[0]["campaign_name"] == "Saisie"


def test_light_yellow_counts_and_green_does_not(tmp_path):
    fills = {(1, 2): LIGHT_YELLOW, (2, 2): LIGHT_YELLOW}
    fills.update(_yellow_rows([6], fill=LIGHT_YELLOW))
    fills.update(_yellow_rows([8], fill=GREEN))
    rows = load_rows_from_path(_book(tmp_path, ROWS, fills))
    assert [r["placement_name"] for r in rows] == ["Article"]


def test_without_yellow_the_whole_file_is_read(tmp_path):
    rows = load_rows_from_path(_book(tmp_path, ROWS, {}))
    assert len(rows) == 4
    assert not any("Mode cases jaunes" in l for l in rows.ingest_report)


def test_merged_yellow_cell_covers_its_range(tmp_path):
    fills = {(1, 2): YELLOW, (2, 2): YELLOW}
    fills.update(_yellow_rows([5, 6], cols=range(2, 7)))
    fills[(5, 1)] = YELLOW
    p = _book(tmp_path, ROWS, fills)
    wb = openpyxl.load_workbook(p)
    wb.active.merge_cells("A5:A6")                               # « Le Monde » sur 2 lignes
    wb.save(p)
    rows = load_rows_from_path(p)
    assert [r["site_name"] for r in rows] == ["Le Monde", "Le Monde"]


@pytest.mark.parametrize("rgb,yes", [("FFFF00", True), ("FFF2CC", True), ("FFEB9C", True),
                                     ("FFC000", True), ("FFFFFF", False), ("C6EFCE", False),
                                     ("FF0000", False), ("BDD7EE", False)])
def test_is_yellow(rgb, yes):
    assert plan_agent.is_yellow(rgb) is yes


def test_theme_colors_resolved(tmp_path):
    """Couleurs de thème (jaune clair Excel = accent « or » éclairci à 80 %)."""
    from openpyxl.styles.colors import Color
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    wb.active["A1"].fill = PatternFill("solid", fgColor=Color(theme=7, tint=0.8))
    wb.save(tmp_path / "t.xlsx")
    wb2 = openpyxl.load_workbook(tmp_path / "t.xlsx")
    theme = plan_agent._theme_colors(wb2)
    assert len(theme) == 12 and theme[0] == "FFFFFF"
    assert plan_agent._fill_rgb(wb2.active["A1"], theme) == plan_agent._apply_tint(theme[7], 0.8)
    assert plan_agent._apply_tint("FFC000", 0.8) == "FFF2CC"      # thème Office 2013+ : jaune clair
