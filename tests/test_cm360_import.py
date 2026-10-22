"""Import du Campaign spreadsheet dans CM360 : sans site, CM360 ne crée que les créatives
(constaté sur un vrai compte). Umnia le signale avant l'export et exporte l'ID de site saisi."""
import json
import os
import sys

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path[:0] = [ROOT, os.path.join(ROOT, "web"), os.path.join(ROOT, "tests", "fixtures")]

pytest.importorskip("openpyxl")


def _analyze(tmp_path, instructions=""):
    from make_upsa import build
    import web_api
    path = tmp_path / "upsa.xlsx"
    build(str(path))
    return web_api, json.loads(web_api.analyze(str(path), "", "", "", "", instructions))


def test_missing_site_is_blocking_for_cm360(tmp_path):
    _, d = _analyze(tmp_path)
    block = [i for i in d["cm360_issues"] if i["key"] == "site" and i["level"] == "bloquant"]
    assert block and block[0]["count"] == 16 and "seules les créatives" in block[0]["message"]
    assert any(i["key"] == "package" for i in d["cm360_issues"])


def test_site_id_exported_and_issue_cleared(tmp_path):
    import openpyxl, io
    web_api, d = _analyze(tmp_path, "Régie : DV360 Open Market")
    assert not [i for i in d["cm360_issues"] if i["key"] == "site"]
    assert any(i["key"] == "site_id" for i in d["cm360_issues"])               # nom seul : alerte
    table = d["table"]
    for r in table:
        r["site_id"] = "8812345"
    d2 = json.loads(web_api.rebuild(json.dumps(table)))
    assert not any(i["key"] in ("site", "site_id") for i in d2["cm360_issues"])
    assert any(c["key"] == "site_id" for c in d2["columns"])
    ws = openpyxl.load_workbook(io.BytesIO(web_api.export_cm360())).active
    head = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]
    assert row[head.index("Site ID\n(Placement)")] == "8812345"
    assert row[head.index("Site name\n(Placement)")] == "DV360 Open Market"
