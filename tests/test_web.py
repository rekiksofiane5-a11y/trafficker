"""Version 100 % navigateur : le moteur Python tourne dans Chromium (Pyodide).
Parcours complet : chargement, analyse d'un plan, simulation, rapport, exports.
Ignoré si node_modules/pyodide ou Playwright manquent (sauf TRAFFICKER_REQUIRE_BROWSER=1)."""
import functools
import http.server
import io
import os
import subprocess
import sys
import threading

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIST = os.path.join(ROOT, "web", "dist")
REQUIRED = os.environ.get("TRAFFICKER_REQUIRE_BROWSER") == "1"
CHROMIUM = os.environ.get("CHROMIUM_PATH", "/opt/pw-browsers/chromium")

if not os.path.isdir(os.path.join(ROOT, "node_modules", "pyodide")) and not REQUIRED:
    pytest.skip("node_modules/pyodide absent (npm ci)", allow_module_level=True)
if REQUIRED:
    import playwright.sync_api as sync_api
else:
    sync_api = pytest.importorskip("playwright.sync_api")


@pytest.fixture(scope="module")
def site():
    if os.environ.get("TRAFFICKER_WEB_URL"):          # ex. wrangler dev (en-têtes réels)
        yield os.environ["TRAFFICKER_WEB_URL"]
        return
    subprocess.run([sys.executable, os.path.join(ROOT, "web", "build.py")], check=True)
    handler = functools.partial(http.server.SimpleHTTPRequestHandler, directory=DIST)
    handler.log_message = lambda *a: None
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_address[1]}/app/"           # site vitrine à la racine, application sous /app/
    srv.shutdown()


@pytest.fixture(scope="module")
def context(site):
    """Un navigateur par module ; le cache HTTP (moteur Pyodide) est partagé entre les tests."""
    with sync_api.sync_playwright() as p:
        kw = {"executable_path": os.path.realpath(CHROMIUM)} if os.path.exists(CHROMIUM) else {}
        if hasattr(os, "geteuid") and os.geteuid() == 0:
            kw["args"] = ["--no-sandbox", "--disable-gpu", "--no-zygote"]
        browser = p.chromium.launch(**kw)
        ctx = browser.new_context(accept_downloads=True)
        yield ctx
        browser.close()


@pytest.fixture(scope="module")
def _page(context, site):
    pg = context.new_page()
    pg.errors = []
    pg.on("pageerror", lambda e: pg.errors.append(str(e)))
    pg.goto(site)
    pg.wait_for_selector("text=Moteur prêt", timeout=120000)
    yield pg
    pg.close()


@pytest.fixture
def page(_page):
    """Même page (le moteur Pyodide n'est chargé qu'une fois) remise à zéro avant chaque test : stockage du
    navigateur vidé (tâches, réglages, alertes), fenêtres fermées, mode Trafficking, une tâche neuve."""
    _page.evaluate("""() => {
      for (const d of document.querySelectorAll("dialog[open]")) d.close();
      localStorage.clear(); sessionStorage.clear();
      closeMenus(); setHistory(false); setHelp(false);
      helpMsgs = []; urlRes.clear(); siteDir = null;
      if (current) { current = null; resetWork(); $("task-view").hidden = true; $("empty-view").hidden = false; }
      $("adserver").value = "cm360";
      setMode("trafficking");
      refreshNotifs();
    }""")
    _page.errors.clear()
    _create_task(_page, "Tâche de test")
    yield _page


def _history(page):
    """L'historique est une fenêtre fermée par défaut : l'ouvrir avant d'y cliquer."""
    if not page.is_visible("#history"):
        page.click("#history-btn")
        page.wait_for_selector("#history:not([hidden])")


def _open_task(page, name):
    _history(page)
    page.click(f".task-item:has-text('{name}')")


def _create_task(page, name, **fields):
    page.evaluate("openCreate()")                          # plus de bouton dans l'historique
    page.fill("#cd-name", name)
    for key, value in fields.items():
        page.fill("#cd-" + key, value)
    page.click("#cd-submit")
    page.wait_for_selector("#task-view:not([hidden])")


def _report_space(page, advertiser):
    """Accueil Reporting : annonceur choisi dans la liste -> « Espace report »."""
    if not page.is_visible("#rp-home"):
        page.click(".modes [data-mode=reporting]")
    page.wait_for_function("document.querySelectorAll('#rp-home-adv option').length > 1")
    label = next(o for o in page.locator("#rp-home-adv option").all_text_contents() if advertiser in o)
    page.select_option("#rp-home-adv", label=label)
    page.click("#rp-home-go")
    page.wait_for_selector("#rp-template:not([hidden])")


def _scope(page, campaign):
    """Après le template : campagne puis dates (pré-remplies), Valider -> la section se réduit."""
    page.wait_for_function("document.querySelectorAll('#rp-camp-sel option').length > 1")
    label = next(o for o in page.locator("#rp-camp-sel option").all_text_contents() if campaign in o)
    page.select_option("#rp-camp-sel", label=label)
    page.click("#rp-validate")
    page.wait_for_selector("#rp-scope-compact:not([hidden])")


def _analyze(page, path, validate=True):
    if page.is_visible("#pm-reopen"):                     # zone de dépôt repliée après une production
        page.click("#pm-reopen")
    page.set_input_files("#file", path)
    page.click("#analyze")
    page.wait_for_selector("#result:not(.hidden)", timeout=30000)
    page.wait_for_function("!document.getElementById('analyze').disabled", timeout=30000)
    if validate and page.is_visible("#validate-step"):      # étape 1 validée : annonces et créas
        page.click("#validate-step")


def test_template_plan_end_to_end(page):
    _analyze(page, os.path.join(ROOT, "plan_media_template.csv"))
    text = page.text_content("#result")
    assert "0 erreur(s)" in text and "Printemps_2026_FR" in text
    assert "Aperçu dans CM360" in text
    assert "Tracking" in text and "In-stream video" in text and "Display" in text
    assert "Placement" in text and "nouveau" in text
    page.click("#demo")
    page.wait_for_selector("text=Trafficking simulé")
    assert "0 erreur(s)" in page.inner_text("#out")
    page.click("#delivery")
    page.wait_for_selector("text=Rapport de diffusion simulé")
    assert "BN_Home_300x250" in page.inner_text("#out") and "€" in page.inner_text("#out")
    assert page.errors == []


@pytest.mark.parametrize("btn,name,sheet", [("#dl-cm360", "campaign_spreadsheet_cm360.xlsx", "Campaign spreadsheet"),
                                            ("#dl-adform", "adform_bulk.xlsx", "Campaigns")])
def test_exports_download_valid_xlsx(page, tmp_path, btn, name, sheet):
    openpyxl = pytest.importorskip("openpyxl")
    _analyze(page, os.path.join(ROOT, "plan_media_template.csv"))
    page.select_option("#adserver", "adform" if "adform" in btn else "cm360")   # choix dans la carte « prêt »
    with page.expect_download() as dl:
        page.click(btn)
    page.select_option("#adserver", "cm360")
    assert dl.value.suggested_filename == name
    out = tmp_path / name
    dl.value.save_as(out)
    assert sheet in openpyxl.load_workbook(out).sheetnames


def test_free_form_plan_and_errors(page, tmp_path):
    p = tmp_path / "libre.csv"
    p.write_text("Campagne,Emplacement,Format,Début,Fin,Annonce,Créatif,URL\n"
                 "C1,P1,300 x 250,02/11/2026,01/11/2026,A1,K1,https://x.fr\n", encoding="utf-8")
    _analyze(page, str(p))
    text = page.text_content("#result")
    assert "campaign_name ← colonne « Campagne »" in text
    assert "à compléter" in text                               # fin avant début : vidée, à compléter
    assert not page.is_disabled("#dl-cm360")                   # ça avance : export possible


def test_nothing_sent_over_network(page, site):
    """Aucune requête vers un autre domaine que le site lui-même."""
    seen = []
    page.on("request", lambda r: seen.append(r.url))
    _analyze(page, os.path.join(ROOT, "plan_media_template.csv"))
    page.click("#demo")
    assert all(u.startswith(site) or u.startswith("blob:") or u.startswith("data:") for u in seen), seen


def test_creative_files_checked_against_cm360_specs(page, tmp_path):
    from test_creative_specs import html5_zip, png
    (tmp_path / "ban.png").write_bytes(png(600, 500))
    (tmp_path / "anim.zip").write_bytes(html5_zip("<html>pas de click</html>"))
    plan = tmp_path / "plan.csv"
    plan.write_text("campaign_name,campaign_start,campaign_end,placement_name,size,ad_name,"
                    "creative_name,creative_path,landing_page_url\n"
                    "C,2026-11-02,2026-12-15,P1,300x250,A1,K1,ban.png,https://x.fr\n"
                    "C,2026-11-02,2026-12-15,P2,300x250,A2,K2,anim.zip,https://x.fr\n", encoding="utf-8")
    page.set_input_files("#creas", [str(tmp_path / "ban.png"), str(tmp_path / "anim.zip")])
    _analyze(page, str(plan))
    text = page.text_content("#result")
    assert "Contrôle des créatives (specs CM360)" in text
    assert "ban.png : conforme" in text and "aucun click tag" in text
    assert "HTML5" in text
    assert page.is_disabled("#dl-cm360")              # créative hors specs : export bloqué
    page.set_input_files("#creas", [])


def test_yellow_cells_mode_in_browser(page, tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.styles import PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in [["Campagne", "Emplacement", "Format", "Début", "Fin", "URL"],
              ["Hiver", "Home", "300x250", "2026-11-02", "2026-12-15", "https://x.fr"],
              ["Hiver", "Article", "728x90", "2026-11-02", "2026-12-15", "https://x.fr"]]:
        ws.append(r)
    for c in range(1, 7):
        ws.cell(row=3, column=c).fill = PatternFill("solid", fgColor="FFFF00")
    wb.save(tmp_path / "jaune.xlsx")
    _analyze(page, str(tmp_path / "jaune.xlsx"))
    text = page.text_content("#result")
    assert "Mode cases jaunes" in text and "Article" in text
    assert "Home" not in page.inner_text("#preview-pane")


def test_page_stays_responsive_on_a_heavy_workbook(page, tmp_path):
    """Gros classeur (feuilles annexes, milliers de lignes) : l'analyse tourne dans le
    Web Worker, la page répond pendant ce temps."""
    import time
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plan"
    ws.append(["Campagne", "Emplacement", "Format", "Début", "Fin", "URL"])
    for i in range(3000):
        ws.append(["Hiver", f"Emplacement {i}", "300x250", "2026-11-02", "2026-12-15", "https://x.fr"])
    annexe = wb.create_sheet("Référentiel")
    for i in range(8000):
        annexe.append([f"code {i}", i, "valeur"])
    wb.save(tmp_path / "lourd.xlsx")
    page.set_input_files("#file", str(tmp_path / "lourd.xlsx"))
    page.click("#analyze")
    t0 = time.time()
    assert page.evaluate("1 + 1") == 2                       # le fil principal répond
    assert time.time() - t0 < 2
    page.wait_for_function("!document.getElementById('analyze').disabled", timeout=120000)
    assert "3000 ligne(s)" in page.text_content("#result")


def test_missing_info_left_empty_and_listed(page, tmp_path):
    p = tmp_path / "incomplet.csv"
    p.write_text("Campagne,Emplacement,Format\nHiver,Home,Pavé\nHiver,Article,Méga-bannière\n", encoding="utf-8")
    _analyze(page, str(p))
    text = page.text_content("#result")
    assert "information(s) à compléter" in text                 # un seul rappel, pas de liste
    assert "Date de début — à compléter" not in text
    assert page.locator("#cf-campaign_start.cell-missing").count() == 1
    assert "erreur(s)" in text and "0 erreur(s)" in text
    assert not page.is_disabled("#dl-cm360")                     # export possible, cases vides
    page.click("#delivery")
    page.wait_for_selector("text=Dates de campagne à compléter")


def test_split_screen_edit_updates_preview(page, tmp_path):
    p = tmp_path / "edit.csv"
    p.write_text("Campagne,Régie,Emplacement,Format,Début,Fin,URL\n"
                 "Hiver,Le Monde,Home,Pavé,02/11/2026,15/12/2026,https://x.fr\n", encoding="utf-8")
    _analyze(page, str(p))
    assert page.is_visible("#editor table") and page.is_visible("#preview-pane .preview")
    assert "Home" in page.inner_text("#preview-pane")
    # l'objectif saisi dans le tableau entre dans la nomenclature de l'aperçu
    page.fill('#editor input[data-key="objective"]', "Notoriété")
    page.wait_for_function("document.getElementById('preview-pane').innerText.includes('NOTORIETE-LE_MONDE-PAVE')",
                           timeout=15000)
    # vider l'URL : case orange + « à compléter »
    page.fill('#editor input[data-key="landing_page_url"]', "")
    page.wait_for_selector("text=information(s) à compléter", timeout=15000)
    assert page.locator("#editor td.cell-missing").count() >= 1
    # ajouter une ligne : nouvel emplacement dans l'aperçu
    page.click("#add-row")
    page.fill('#editor tr[data-row="1"] input[data-key="placement_name"]', "Article")
    page.fill('#editor tr[data-row="1"] input[data-key="format_label"]', "Méga-bannière")
    page.wait_for_function("document.getElementById('preview-pane').innerText.includes('728 × 90')", timeout=15000)


def test_campaign_header_applies_to_all_rows(page, tmp_path):
    p = tmp_path / "entete.csv"
    p.write_text("Objectif,Régie,Emplacement,Format,URL\n"
                 "Notoriété,Le Monde,Home,Pavé,https://x.fr\n"
                 "Trafic,Le Monde,Article,Pavé,https://x.fr\n", encoding="utf-8")
    _analyze(page, str(p))
    assert page.is_visible("#cf-campaign_name")
    page.fill("#cf-campaign_name", "Printemps 2027")
    page.fill("#cf-campaign_start", "01/03/2027")
    page.fill("#cf-campaign_end", "31/03/2027")
    page.wait_for_function("document.getElementById('preview-pane').innerText.includes('Printemps 2027')", timeout=15000)
    pane = page.inner_text("#preview-pane")
    assert "01/03/2027" in pane and pane.count("Printemps 2027") == 1        # une seule campagne
    assert "PACKAGE-LE_MONDE-NOTORIETE" in pane and "PACKAGE-LE_MONDE-TRAFIC" in pane


def test_adform_preview_with_views(page, tmp_path):
    p = tmp_path / "adform.csv"
    p.write_text("Campagne,Objectif,Régie,Emplacement,Format,Début,Fin,Budget net,Volume (imp.),URL\n"
                 "SPF Fertilité,Notoriété,Amnet,Contextuel Santé,Audio,21/10/2026,22/11/2027,151619,11544211,https://x.fr\n"
                 "SPF Fertilité,Notoriété,Amnet,Comportemental,Audio,21/10/2026,22/11/2027,151619,11544211,https://x.fr\n",
                 encoding="utf-8")
    _analyze(page, str(p))
    page.select_option("#adserver", "adform")                # choix dans la carte « prêt »
    page.wait_for_function("document.getElementById('preview-pane').innerText.includes('Aperçu dans Adform')")
    pane = page.inner_text("#preview-pane")
    assert "Aperçu dans Adform" in pane and "Direct Line Item" in pane and "Amnet" in pane
    assert "151\u202f619,00 €" in pane and "Imps" in pane and "21/10/2026" in pane
    assert page.is_visible("#dl-adform") and not page.is_visible("#dl-cm360")
    page.select_option("#adform-view", "banners")
    assert "Line item" in page.inner_text("#preview-pane")
    page.select_option("#adform-view", "tags")
    assert "Tag type" in page.inner_text("#preview-pane")
    page.select_option("#adserver", "cm360")                    # retour CM360 : arborescence
    assert "Aperçu dans CM360" in page.inner_text("#preview-pane")


def test_step_by_step_and_direct_edit_in_preview(page, tmp_path):
    p = tmp_path / "etapes.csv"
    p.write_text("Campagne,Objectif,Régie,Emplacement,Format,Message,Début,Fin,URL\n"
                 "Hiver,Notoriété,Le Monde,Home,Pavé,Soldes,02/11/2026,15/12/2026,https://x.fr\n"
                 "Hiver,Notoriété,Le Monde,Home,Pavé,Noël,02/11/2026,15/12/2026,https://x.fr\n"
                 "Hiver,Trafic,L'Équipe,Article,Méga-bannière,Soldes,02/11/2026,15/12/2026,https://x.fr\n",
                 encoding="utf-8")
    _analyze(page, str(p), validate=False)
    page.select_option("#adserver", "cm360")
    pane = lambda: page.inner_text("#preview-pane")
    # Étape 1 : emplacements sous la régie, pas encore d'annonces ni de créas
    assert "1. Emplacements" in pane() and "Placement" in pane() and "Tracking" not in pane()
    assert page.is_disabled("#dl-cm360")                               # export après validation
    # Grouper en package : la colonne Package du plan est remplie
    page.check('#preview-pane tr.pv-placement input[type=checkbox] >> nth=0')
    page.once("dialog", lambda d: d.accept("Package-Test"))
    page.click("#group-package")
    page.wait_for_function("document.getElementById('preview-pane').innerText.includes('Package-Test')", timeout=15000)
    assert page.locator('#editor input[data-key="package"]').first.input_value() == "Package-Test"
    # Renommer un emplacement par double-clic : reporté dans le plan (colonne imposée)
    page.locator("#preview-pane tr.pv-placement .pv-link").first.dblclick()
    page.fill("input.pv-rename", "MON-EMPLACEMENT")
    page.keyboard.press("Enter")
    page.wait_for_function("document.getElementById('preview-pane').innerText.includes('MON-EMPLACEMENT')", timeout=15000)
    assert page.locator('#editor input[data-key="placement_final"]').count() >= 1
    assert page.locator("#editor tr.hl").count() >= 1                  # lignes du plan surlignées
    # Glisser l'autre emplacement sur le site « Le Monde »
    src = page.locator("#preview-pane tr.pv-placement").nth(1)
    dst = page.locator("#preview-pane tr.pv-site", has_text="Le Monde")
    src.drag_to(dst)
    page.wait_for_function("document.querySelectorAll('#preview-pane tr.pv-site').length === 1", timeout=15000)
    # Étape 2 : annonces = emplacement × créative, nom par défaut « emplacement-créative »
    page.click("#validate-step")
    text = pane()
    assert "2. Annonces & créatives" in text and "MON-EMPLACEMENT-" in text
    assert not page.is_disabled("#dl-cm360")
    ad = page.locator("#preview-pane tr.pv-ad .pv-link").first
    ad.dblclick()
    page.fill("input.pv-rename", "ANNONCE-PERSO")
    page.keyboard.press("Enter")
    page.wait_for_function("document.getElementById('preview-pane').innerText.includes('ANNONCE-PERSO')", timeout=15000)
    assert page.locator('#editor input[data-key="ad_final"]').count() >= 1


def test_tasks_are_saved_restored_and_reporting_mode(page):
    """Coquille : tâche créée, production lancée, rechargement -> la tâche se rouvre telle quelle."""
    _create_task(page, "Lancement printemps", advertiser="Annonceur Démo", campaign="Printemps")
    page.fill("#instructions", "URL : https://www.exemple.fr")
    _analyze(page, os.path.join(ROOT, "plan_media_template.csv"))
    page.wait_for_function("document.getElementById('task-chips').textContent.includes('emplacements validés')")
    page.reload()
    page.wait_for_selector("text=Moteur prêt", timeout=120000)
    assert page.is_visible("#empty-view")
    _open_task(page, "Lancement printemps")
    page.wait_for_selector("#result:not(.hidden)")
    page.wait_for_selector("#back-step")                       # étape 2 restaurée
    assert page.input_value("#cf-campaign_name") == "Printemps"      # la pop-up fait foi
    assert page.input_value("#cf-advertiser_name") == "Annonceur Démo"
    assert page.input_value("#instructions") == "URL : https://www.exemple.fr"
    _history(page)
    page.click(".scopes [data-scope=agency]")
    page.wait_for_selector("#task-list .task-item")
    page.click(".modes [data-mode=reporting]")
    assert page.is_visible("#empty-view") and page.is_visible("#rp-home") and not page.is_visible("#empty-new")
    _report_space(page, "UPSA")
    assert page.is_visible("#rp-view") and not page.is_visible("#pm-step")
    assert "Report UPSA" in page.text_content("#task-title") or "UPSA" in page.text_content("#rp-adv-label")
    page.click(".modes [data-mode=trafficking]")
    assert page.errors == []


def test_resizable_columns_and_folded_drop_zone(page, tmp_path):
    from fixtures.make_upsa import build
    path = tmp_path / "upsa.xlsx"
    build(str(path))
    _create_task(page, "UPSA Citrate")
    _analyze(page, str(path), validate=False)
    assert page.is_visible("#pm-compact") and not page.is_visible("#drop")      # dépôt replié
    assert "upsa.xlsx" in page.text_content("#pm-compact-text")
    th = page.locator("#editor th", has_text="Détail du ciblage")
    w0 = th.bounding_box()["width"]
    grip = th.locator(".col-grip")
    grip.scroll_into_view_if_needed()
    box = grip.bounding_box()
    page.mouse.move(box["x"] + 4, box["y"] + 5)
    page.mouse.down()
    page.mouse.move(box["x"] + 154, box["y"] + 5, steps=5)
    page.mouse.up()
    assert th.bounding_box()["width"] > w0 + 100                                # colonne agrandie
    assert "CONTEXTUEL" in page.text_content("#preview-pane")                # un package par ciblage
    page.click("#pm-reopen")
    assert page.is_visible("#drop")
    assert page.errors == []


def test_bulk_edit_modify_duplicate_delete(page, tmp_path):
    from fixtures.make_upsa import build
    path = tmp_path / "upsa.xlsx"
    build(str(path))
    _create_task(page, "UPSA bulk")
    _analyze(page, str(path), validate=False)
    rows = lambda: page.locator("#editor tbody tr").count()
    assert rows() == 16 and not page.is_visible("#bulk-bar")
    page.click('#editor tr[data-row="0"] td.sel input')
    page.click('#editor tr[data-row="3"] td.sel input', modifiers=["Shift"])         # plage 1 à 4
    assert page.text_content("#bulk-count") == "4 ligne(s) sélectionnée(s)"
    page.select_option("#bulk-key", "site_name")
    page.fill("#bulk-value", "DV360")
    page.click("#bulk-apply")
    page.wait_for_function("document.getElementById('preview-pane').innerText.includes('DV360')", timeout=15000)
    values = [page.input_value(f'#editor tr[data-row="{i}"] input[data-key="site_name"]') for i in range(5)]
    assert values == ["DV360"] * 4 + [""]
    page.click("#bulk-dup")
    assert rows() == 20 and page.text_content("#bulk-count") == "4 ligne(s) sélectionnée(s)"   # copies sélectionnées
    page.once("dialog", lambda d: d.accept())
    page.click("#bulk-delete")
    assert rows() == 16 and not page.is_visible("#bulk-bar")
    page.check("#ed-all")
    assert page.text_content("#bulk-count") == "16 ligne(s) sélectionnée(s)"
    page.click("#bulk-clear")
    assert not page.is_visible("#bulk-bar")
    assert page.errors == []


def test_reporting_check_and_scheduled_reports(page):
    """Reporting (démo, sans ad server connecté) : annonceur -> campagne -> check de mise en ligne
    (tableau de bord + Excel) puis report d'envoi par régie (carnet d'adresses, aperçu, programmation)."""
    openpyxl = pytest.importorskip("openpyxl")
    page.click(".modes [data-mode=reporting]")
    assert "démonstration" in page.text_content("#rp-home-src")
    _report_space(page, "Automobiles Démo")
    assert "démonstration" in page.text_content("#rp-source")
    page.click(".tpl[data-tpl=check]")
    assert not page.is_visible(".kpis")                            # tableau de bord après campagne et dates
    _scope(page, "Lancement SUV")
    page.wait_for_selector(".kpis")
    assert "4 régie(s)" in page.text_content("#rp-selected") and "Lancement SUV" in page.text_content("#rp-scope-text")
    dash = page.text_content("#rp-dash")
    assert "Régies avec données" in dash and "Aucune donnée" in dash and "Données reçues" in dash
    assert page.locator("#rp-check-table tbody tr").count() == 7
    with page.expect_download() as dl:
        page.click("#rp-dl-check")
    assert dl.value.suggested_filename == "check_mise_en_ligne_Lancement_SUV_electrique.xlsx"
    assert openpyxl.load_workbook(io.BytesIO(open(dl.value.path(), "rb").read())).sheetnames == ["Synthèse", "Par régie", "Emplacements", "Données jour"]
    page.click(".tpl[data-tpl=regies]")
    page.wait_for_selector(".rp-regie")
    options = page.locator(".rp-regie select option").all_text_contents()
    assert len(options) == 4 and all(any(n in o for n in ("Le Monde", "Les Échos", "Prisma", "Teads")) for o in options)
    page.click("#rp-launch")
    assert "Ajoutez au moins un e-mail" in page.text_content("#rp-send-msg")
    page.fill(".rp-regie input", "trafic@lemonde.fr")
    page.click("#rp-add-regie")
    assert page.locator(".rp-regie").count() == 2
    page.locator(".rp-regie").nth(1).locator("input").fill("pas-un-mail")
    page.click("#rp-launch")
    assert "E-mail invalide" in page.text_content("#rp-send-msg")
    page.locator(".rp-regie").nth(1).locator("input").fill("adops@regie.fr")
    page.locator(".rp-regie").nth(0).locator("button:has-text('Prévisualiser')").click()
    page.wait_for_selector("#rp-preview:not([hidden])")
    assert "Total" in page.text_content("#rp-preview") and "Impressions" in page.text_content("#rp-preview")
    page.click("#rp-launch")
    page.wait_for_selector("#rp-send-msg.ok")
    assert page.locator("#rp-scheduled tbody tr").count() == 2
    assert "trafic@lemonde.fr" in page.text_content(".rp-regie .sugg")                # carnet d'adresses
    _open_task(page, "Report Automobiles Démo")                     # réouverture
    page.wait_for_selector("#rp-template:not([hidden])")
    assert page.locator(".tpl.on").get_attribute("data-tpl") == "regies"
    page.click(".modes [data-mode=trafficking]")
    assert page.errors == []


def test_pacing_from_linked_plan_and_manage_scheduled_reports(page, tmp_path):
    """Check : rythme prévu (plan de la tâche de trafficking liée) vs livré ; reports programmés :
    pause, reprise, modification, suppression."""
    from fixtures.make_upsa import build
    path = tmp_path / "upsa.xlsx"
    build(str(path), 1600000, "8 000 €")
    _create_task(page, "UPSA Citrate — trafficking", advertiser="UPSA")
    page.fill("#instructions", "Régie : DV360 Open Market")       # noms finaux = ceux de la campagne (démo)
    _analyze(page, str(path), validate=False)
    page.wait_for_function("document.getElementById('task-chips').textContent.includes('produit')")
    _report_space(page, "UPSA")
    page.click(".tpl[data-tpl=check]")
    _scope(page, "Citrate")
    page.wait_for_selector("#rp-pace-table")
    assert "UPSA Citrate — trafficking" in page.locator("#rp-plan option:checked").text_content()   # proposé seul
    dash = page.text_content("#rp-dash")
    assert "Rythme : livré / prévu à date" in dash and "Budget consommé (estimé)" in dash
    assert "16 emplacement(s) rapprochés du plan" in dash
    assert "sur 8 000 €" in dash.replace("\u202f", " ").replace("\xa0", " ")
    page.select_option("#rp-plan", "")
    page.wait_for_function("!document.getElementById('rp-pace-table')")
    page.click(".tpl[data-tpl=regies]")
    page.fill(".rp-regie input", "trafic@regie.fr")
    page.click("#rp-launch")
    page.wait_for_selector("#rp-sched-table")
    page.click("#rp-sched-table button:has-text('Pause')")
    assert "En pause" in page.text_content("#rp-sched-table") and "mis en pause" in page.text_content("#rp-send-msg")
    page.click("#rp-sched-table button:has-text('Reprendre')")
    assert "Actif" in page.text_content("#rp-sched-table")
    page.click("#rp-sched-table button:has-text('Modifier')")
    page.fill("#se-emails", "trafic@regie.fr, ops@regie.fr")
    page.select_option("#se-freq", "daily")
    page.click("#se-save")
    row = page.text_content("#rp-sched-table tbody tr")
    assert "ops@regie.fr" in row and "Quotidien" in row
    page.once("dialog", lambda d: d.accept())
    page.click("#rp-sched-table button:has-text('Supprimer')")
    assert not page.is_visible("#rp-sched-table") and "supprimé" in page.text_content("#rp-send-msg")
    page.click(".modes [data-mode=trafficking]")
    assert page.errors == []


def test_cm360_prompt_without_api(page):
    """Sans accès API : demande en français -> questions -> prompt complet pour « Generate report »."""
    page.click(".modes [data-mode=reporting]")
    _create_task(page, "Prompt visibilité")
    assert not page.is_disabled(".tpl[data-tpl=prompt]")
    page.click(".tpl[data-tpl=prompt]")
    page.fill("#pp-input", "visibilité")
    page.click("#pp-go")
    page.wait_for_selector("#pp-answer")
    page.fill('[data-q="campaign"] input', "Noël 2026")
    page.click("#pp-answer")
    page.wait_for_selector("#pp-prompt")
    prompt = page.input_value("#pp-prompt")
    assert 'campaign "Noël 2026"' in prompt and "for last week" in prompt and "Active View" in prompt
    assert "Placement ID" in page.text_content("#pp-settings")
    page.click("#pp-copy")
    _open_task(page, "Prompt visibilité")                   # réouverture : prompt restauré
    page.wait_for_selector("#pp-prompt")
    assert page.input_value("#pp-prompt") == prompt
    page.click("#pp-presets button:has-text('Report régie')")          # sans campagne choisie : questions
    page.fill('[data-q="campaign"] input', "Noël 2026")
    page.click("#pp-answer")
    page.wait_for_selector("#pp-prompt")
    assert "Rows (dimensions): Campaign, Campaign ID, Date" in page.input_value("#pp-prompt")
    assert "Schedule and delivery" in page.text_content("#pp-output")
    page.click(".modes [data-mode=trafficking]")
    assert page.errors == []


def test_reporting_from_cm360_export_file(page, tmp_path):
    """Sans API : export de rapport CM360 déposé -> check (avec plan lié : emplacements muets) et report régie."""
    from fixtures.make_upsa import build
    plan = tmp_path / "upsa.xlsx"
    build(str(plan), 1600000, "8 000 €")
    _create_task(page, "UPSA plan (fichier)", advertiser="UPSA")
    page.fill("#instructions", "Régie : DV360 Open Market")
    _analyze(page, str(plan), validate=False)
    page.wait_for_function("document.getElementById('task-chips').textContent.includes('produit')")
    page.click(".modes [data-mode=reporting]")
    _create_task(page, "UPSA export")
    page.set_input_files("#rp-file", os.path.join(ROOT, "web", "exemple_export_cm360.csv"))
    page.wait_for_selector("#rp-fstart")
    assert "72 ligne(s)" in page.text_content("#rp-selected") and not page.is_disabled(".tpl[data-tpl=check]")
    page.fill("#rp-fend", "2026-10-23")
    page.dispatch_event("#rp-fend", "change")
    page.click(".tpl[data-tpl=check]")
    page.wait_for_selector("#rp-pace-table")
    dash = page.text_content("#rp-dash")
    assert "12 / 16" in dash                                    # 4 emplacements du plan absents de l'export
    assert "Données de l'export" in page.text_content("#rp-check-status")
    page.click(".tpl[data-tpl=regies]")
    page.fill(".rp-regie input", "trafic@regie.fr")
    page.click("#rp-launch")
    page.wait_for_selector("#rp-sched-table")
    assert "à programmer dans CM360" in page.text_content("#rp-sched-table")
    _open_task(page, "UPSA export")             # réouverture : export conservé
    page.wait_for_selector("#rp-fstart")
    assert page.input_value("#rp-fend") == "2026-10-23"
    page.click(".modes [data-mode=trafficking]")
    assert page.errors == []


def test_real_cm360_xlsx_export_and_keyword_filter(page, tmp_path):
    """Excel exporté de CM360 (non standard) : import sans erreur, alerte d'export trop large, filtre par mot."""
    from fixtures.make_cm360_export import build
    path = tmp_path / "3021_query.xlsx"
    build(str(path))
    page.click(".modes [data-mode=reporting]")
    _create_task(page, "Lotus check")
    page.set_input_files("#rp-file", str(path))
    page.wait_for_selector("#rp-fkw")
    sel = page.text_content("#rp-selected")
    assert "Import impossible" not in sel and "Compte CM360 3021" in sel and "Pas de colonne Campaign" in sel
    page.fill("#rp-fkw", "LEMONDE")
    page.click("#rp-fkw-go")
    page.wait_for_function("document.getElementById('rp-selected').textContent.includes('Filtre « LEMONDE »')")
    assert "1 emplacement(s) avec données" in page.text_content("#rp-selected")
    page.click(".tpl[data-tpl=check]")
    page.wait_for_selector(".kpis")
    assert "179 000" in page.text_content("#rp-dash").replace(" ", " ")
    page.click(".modes [data-mode=trafficking]")
    assert page.errors == []


def test_app_bar_history_and_menus(page):
    """Barre du haut (modes, recherche, thème, compte, historique) et panneau Historique (récents)."""
    page.click(".modes [data-mode=trafficking]")
    _create_task(page, "Tâche historique")
    assert not page.is_visible("#history")                                   # fermé à l'ouverture de la tâche
    main_w = page.locator("#main").bounding_box()["width"]
    _history(page)
    assert page.locator("#main").bounding_box()["width"] == main_w            # ne rétrécit pas la zone de travail
    bg = page.evaluate("getComputedStyle(document.getElementById('history')).backgroundColor")
    assert bg.startswith("rgba(") or "/ 0.7" in bg, bg                         # transparente à 30 %
    item = page.locator(".task-item:has-text('Tâche historique')")
    assert "Brouillon" in item.text_content() and "instant" in item.text_content()
    assert page.locator("#search").count() == 0                              # plus de recherche en haut
    page.fill("#history-q", "historique")
    page.wait_for_function("document.querySelectorAll('#task-list .task-item').length === 1")
    page.fill("#history-q", "")
    page.dispatch_event("#history-q", "input")
    box, main = page.locator("#history").bounding_box(), page.locator("#main").bounding_box()
    assert box["x"] > main["x"]                                                 # historique à droite
    page.click("#history-close")
    assert not page.is_visible("#history") and page.get_attribute("#history-btn", "aria-expanded") == "false"
    page.click("#history-btn")
    assert page.is_visible("#history")
    page.keyboard.press("Escape")                                              # Échap ferme
    assert not page.is_visible("#history")
    _history(page)
    page.mouse.click(40, 400)                                                  # clic à côté ferme
    assert not page.is_visible("#history")
    assert page.text_content("#theme-btn").strip() == ""                       # icône seule
    before = page.evaluate("document.documentElement.classList.contains('is-dark')")
    page.click("#theme-btn")
    assert page.evaluate("document.documentElement.classList.contains('is-dark')") != before
    page.click("#theme-btn")
    assert page.evaluate("document.documentElement.classList.contains('is-dark')") == before
    page.click("#account-btn")
    assert page.is_visible("#open-settings")
    page.keyboard.press("Escape")
    assert not page.is_visible("#account-menu")
    assert page.errors == []


def test_pivot_table_after_export_import(page):
    """Export importé : tableau croisé par défaut (régie), préréglage croisé, tri, Excel."""
    openpyxl = pytest.importorskip("openpyxl")
    page.click(".modes [data-mode=reporting]")
    _create_task(page, "Croisé export")
    page.set_input_files("#rp-file", os.path.join(ROOT, "web", "exemple_export_cm360.csv"))
    page.wait_for_selector("#pv-table")                                      # vue par défaut après import
    head = page.locator("#pv-table thead th").all_text_contents()
    assert head[0] == "Site (CM360)" and "Impressions" in head and "Click Rate" in head
    total = page.text_content("#pv-table tr.total").replace(" ", " ")
    assert "1 747 655" in total and "0,23 %" in total                         # CTR recalculé sur le total
    page.click("#pv-presets button:has-text('Par emplacement')")
    assert page.locator("#pv-table tbody tr:not(.total)").count() == 12
    page.click("#pv-presets button:has-text('Régie × semaine')")
    assert "S43" in " ".join(page.locator("#pv-table thead th").all_text_contents())
    page.select_option("#pv-r1", "Date")
    page.select_option("#pv-col", "")
    page.locator("#pv-table thead th", has_text="Date").click()                # tri croissant par date
    first = page.locator("#pv-table tbody tr").first.text_content()
    assert first.startswith("2026-10-16")
    with page.expect_download() as dl:
        page.click("#pv-dl")
    wb = openpyxl.load_workbook(io.BytesIO(open(dl.value.path(), "rb").read()))
    assert wb.sheetnames == ["Tableau croisé", "Données"]
    _open_task(page, "Croisé export")                       # réouverture : vue conservée
    page.wait_for_selector("#pv-table")
    assert page.input_value("#pv-r1") == "Date"
    page.click(".modes [data-mode=trafficking]")
    assert page.errors == []


def test_nomenclature_dimensions_chart_and_editor_grouping(page):
    """Export à la nomenclature d'agence : dimensions Éditeur / Format, graphique, check regroupé par éditeur."""
    page.click(".modes [data-mode=reporting]")
    _create_task(page, "Lotus nomenclature")
    page.set_input_files("#rp-file", os.path.join(ROOT, "web", "exemple_export_cm360_une_campagne.xlsx"))
    page.wait_for_selector("#pv-table")
    assert "Nomenclature reconnue" in page.text_content("#rp-selected")
    page.click("#pv-presets button:has-text('Par éditeur')")
    assert page.input_value("#pv-r1") == "Éditeur"
    assert page.is_visible("#pv-chart svg") and "Impressions par Éditeur" in page.text_content("#pv-chart")
    assert page.locator("#pv-chart .bar").count() == 8
    page.click("#pv-presets button:has-text('Par jour')")
    assert page.locator("#pv-chart path.line").count() == 1                 # courbe pour une dimension temporelle
    page.click(".tpl[data-tpl=check]")
    page.wait_for_selector(".kpis")
    page.select_option("#rp-group", "editor")
    page.wait_for_function("document.getElementById('rp-dash').textContent.includes('LEMONDE')")
    assert "GRANDANGLE" in page.text_content("#rp-check-table")                # format lu dans la nomenclature
    assert "CPM → impr." in page.text_content("#rp-check-table")               # mode d'achat : impressions vérifiées
    assert "8 CPM · 1 dCPM" in page.text_content(".kpis")
    page.click(".modes [data-mode=trafficking]")
    assert page.errors == []


def test_site_mapping_filters_and_saved_views(page):
    """Correspondance des régies (agence), filtres du croisé, vues enregistrées."""
    page.click(".modes [data-mode=reporting]")
    _create_task(page, "Lotus réglages")
    page.set_input_files("#rp-file", os.path.join(ROOT, "web", "exemple_export_cm360_une_campagne.xlsx"))
    page.wait_for_selector("#pv-table")
    page.click("#rp-sitemap-open")
    page.fill('.sitemap input[data-site="Facebook FR"]', "Meta")
    page.fill('.sitemap input[data-site="Youtube FR"]', "Google")
    page.click("#sitemap-save")
    page.wait_for_function("document.getElementById('rp-sitemap-open').textContent.includes('2 règle')")
    page.click("#pv-presets button:has-text('Par régie')")
    regies = page.locator("#pv-table tbody tr:not(.total) td:first-child").all_text_contents()
    assert "Meta" in regies and "Google" in regies and "Facebook FR" not in regies
    page.click("#pv-add-filter")
    page.select_option("#pvf-dim", "Site (CM360)")
    page.check('#pv-filter-editor input[value="Google"]')
    page.click("#pvf-apply")
    assert "Google" not in page.locator("#pv-table tbody tr:not(.total) td:first-child").all_text_contents()
    assert "sauf Google" in page.text_content("#pv-filter-chips")
    page.once("dialog", lambda d: d.accept("Sans YouTube"))
    page.click("#pv-save-view")
    page.wait_for_selector("#pv-saved button:has-text('Sans YouTube')")
    page.click("#pv-filter-chips button")                                     # retire le filtre
    assert "Google" in page.locator("#pv-table tbody tr:not(.total) td:first-child").all_text_contents()
    page.click("#pv-saved button:has-text('Sans YouTube')")                   # la vue le remet
    assert "Google" not in page.locator("#pv-table tbody tr:not(.total) td:first-child").all_text_contents()
    # écran « Réglages de l'agence » : règles et vues modifiables hors import
    page.click("#account-btn")
    page.click("#open-agency")
    assert page.is_visible("#agency-dialog")
    assert page.locator("#ag-map tbody tr").count() == 2
    page.click('#ag-map button[aria-label="Supprimer la règle Youtube FR"]')
    page.wait_for_function("document.querySelectorAll('#ag-map tbody tr').length === 1")
    page.fill("#ag-map-from", "Teads FR")
    page.fill("#ag-map-to", "Teads")
    page.click("#ag-map-form button[type=submit]")
    page.wait_for_function("document.getElementById('ag-map').textContent.includes('Teads FR')")
    page.fill('#ag-views input[aria-label="Nom de la vue Sans YouTube"]', "Hors Google")
    page.dispatch_event('#ag-views input[aria-label="Nom de la vue Sans YouTube"]', "change")
    page.wait_for_selector("#pv-saved button:has-text('Hors Google')")
    stored = page.evaluate("JSON.parse(localStorage.getItem('umnia_set_site_map'))")
    assert stored == {"Facebook FR": "Meta", "Teads FR": "Teads"}
    page.click("#agency-dialog [data-close]")
    regies = page.locator("#pv-table tbody tr:not(.total) td:first-child").all_text_contents()
    assert "Youtube FR" in regies and "Google" not in regies                  # import réappliqué sans la règle supprimée
    page.evaluate("localStorage.removeItem('umnia_set_site_map'); localStorage.removeItem('umnia_set_pivot_views')")
    page.click(".modes [data-mode=trafficking]")
    assert page.errors == []


def test_nomenclature_labels_and_site_fallback(page, tmp_path):
    """Hors gabarit : éditeur repris du site CM360 ; position renommée dans les réglages, gardée pour l'agence."""
    from fixtures import make_cm360_export
    eds = [("LEMONDE", "GRANDANGLE", "CUSTOM"), ("YOUTUBE", "BUMPER", "SOCIODEMO"), ("AMNET", "NATIVE", "COMPORTEMENTAL"),
           ("TEADS", "INSTREAM", "CUSTOM"), ("LEFIGARO", "PAVE", "SOCIODEMO"), ("M6", "PREROLL", "CUSTOM")]
    data = [("2026-10-01", "DV360", f"R-NOTORIETE-CPM-{e}--{c}-Seg{k}-{f}-1x1-Multidevice-NA", str(100 + k), 1000 * (k + 1), 5, 0.5)
            for k, (e, f, c) in enumerate(eds)]
    data += [("2026-10-01", "Doctolib", "Mon Espace Santé", "900", 4000, 9, 0.22), ("2026-10-02", "Doctolib", "Mon Espace Santé", "900", 3000, 6, 0.2)]
    path = str(tmp_path / "gabarits.xlsx")
    make_cm360_export.build(path, data=data)
    page.click(".modes [data-mode=reporting]")
    _create_task(page, "Lotus gabarits")
    page.set_input_files("#rp-file", path)
    page.wait_for_selector("#pv-table")
    page.click("#pv-presets button:has-text('Par éditeur')")
    eds = page.locator("#pv-table tbody tr:not(.total) td:first-child").all_text_contents()
    assert "(hors nomenclature)" not in eds and "Doctolib (site CM360)" in eds
    page.select_option("#pv-r1", "Gabarit")
    gab = page.locator("#pv-table tbody tr:not(.total) td:first-child").all_text_contents()
    assert "hors nomenclature" in gab and any(g.endswith("segments") for g in gab)
    page.click("#account-btn")
    page.click("#open-agency")
    inp = page.locator("#ag-nom input").first
    learned = inp.get_attribute("placeholder")
    inp.fill("Annonceur ciblé")
    inp.dispatch_event("change")
    page.wait_for_function("(JSON.parse(localStorage.getItem('umnia_set_nom_labels') || '{}'))")
    stored = page.evaluate("JSON.parse(localStorage.getItem('umnia_set_nom_labels'))")
    assert list(stored.values()) == ["Annonceur ciblé"]
    page.click("#agency-dialog [data-close]")
    page.wait_for_function("[...document.querySelectorAll('#pv-r1 option')].some((o) => o.value === 'Annonceur ciblé')", timeout=5000)
    dims = page.locator("#pv-r1 option").all_text_contents()
    assert "Annonceur ciblé" in dims and learned not in dims
    page.evaluate("localStorage.removeItem('umnia_set_nom_labels')")
    page.click(".modes [data-mode=trafficking]")
    assert page.errors == []


def test_pivot_period_comparison(page):
    """Croisé : semaine A contre semaine B, variations signées, dates choisies, Excel."""
    openpyxl = pytest.importorskip("openpyxl")
    page.click(".modes [data-mode=reporting]")
    _create_task(page, "Croisé comparaison")
    page.set_input_files("#rp-file", os.path.join(ROOT, "web", "exemple_export_cm360.csv"))
    page.wait_for_selector("#pv-table")
    page.select_option("#pv-cmp", "week")
    page.wait_for_function("document.querySelector('#pv-table thead').textContent.includes('Impressions · A')")
    head = page.locator("#pv-table thead th").all_text_contents()
    assert head[1:4] == ["Impressions · A", "Impressions · B", "Δ Impressions (%)"]
    assert "Δ Click Rate (pts)" in head
    assert "A : " in page.text_content("#pv-info")
    assert page.locator("#pv-chart rect.mark").count() == page.locator("#pv-chart rect.mark-b").count() >= 1   # barres A et B
    deltas = page.locator("#pv-table tr.total td.delta").all_text_contents()
    assert deltas and all(d == "—" or d[0] in "+−0" or d.endswith(("%", "pt")) for d in deltas)
    page.select_option("#pv-cmp", "custom")
    assert page.is_visible("#pv-cmp-dates") and page.input_value("#pv-a1")     # préremplies
    page.fill("#pv-b1", "2030-01-02")
    page.fill("#pv-b2", "2030-01-01")
    page.dispatch_event("#pv-b2", "change")
    assert "Comparaison impossible" in page.text_content("#pv-info")
    page.select_option("#pv-cmp", "7d")
    page.wait_for_selector("#pv-table")
    with page.expect_download() as dl:
        page.click("#pv-dl")
    wb = openpyxl.load_workbook(io.BytesIO(open(dl.value.path(), "rb").read()))
    ws = wb["Tableau croisé"]
    assert any("Impressions · B" == c.value for row in ws.iter_rows() for c in row)
    page.select_option("#pv-col", "Semaine")                                  # colonnes : fin de la comparaison
    assert page.input_value("#pv-cmp") == ""
    page.click(".modes [data-mode=trafficking]")
    assert page.errors == []


def test_home_import_flow_with_engine_stages(page):
    """Accueil : « Importer mon PM » crée la tâche, puce du fichier, étapes réelles du moteur, « Votre PM est prêt »."""
    page.click(".modes [data-mode=trafficking]")
    page.evaluate("closeTask()")
    assert page.is_visible("#empty-view") and "Un nouveau projet" in page.text_content("#empty-title")
    page.set_input_files("#hero-file", os.path.join(ROOT, "plan_media_template.csv"))
    page.wait_for_selector("#file-chip:not([hidden])")
    assert page.text_content("#task-title") == "plan_media_template"            # tâche nommée d'après le fichier
    assert page.text_content("#fc-name") == "plan_media_template.csv" and not page.is_visible("#drop")
    page.wait_for_function("!document.getElementById('analyze').disabled", timeout=120000)
    page.evaluate("""(() => { window.__stages = []; const o = engine.onmessage;
                     engine.onmessage = (e) => { if (e.data.progress) __stages.push(e.data.progress); o(e); }; })()""")
    page.click("#analyze")
    page.wait_for_selector("#pm-progress:not([hidden])")
    page.wait_for_selector("#pm-compact:not([hidden])", timeout=30000)
    assert page.evaluate("__stages") == ["lecture", "structure", "nommage", "controles", "apercu"]
    assert page.locator("#pm-stages li.done").count() == 5
    assert "Votre PM est prêt" in page.text_content("#pm-compact")
    assert "emplacement(s)" in page.text_content("#pm-compact-text") and page.is_visible("#result")
    assert "emplacements" in page.text_content("#pm-stats") and page.is_visible("#adserver")    # carte « prêt »
    assert page.locator("#checks-list li").count() >= 1                                 # points à vérifier
    assert "politique dans l'UE" in page.text_content("#checks-list")                   # contrôle CM360 remonté
    assert not page.is_visible("#cm360-sheet")                                          # détail replié
    page.click("#pm-reopen")                                                    # importer un autre PM
    assert page.is_visible("#drop") and not page.is_visible("#file-chip")
    page.set_input_files("#file", os.path.join(ROOT, "plan_media_template.csv"))
    page.click("#fc-clear")
    assert page.is_visible("#drop")
    page.click(".modes [data-mode=reporting]")
    page.evaluate("closeTask()")
    assert not page.is_visible("#hero-drop") and not page.is_visible("#empty-new")   # reporting : pas de dépôt de PM
    assert page.is_visible("#rp-home") and page.is_disabled("#rp-home-go")         # annonceur à choisir
    page.click(".modes [data-mode=trafficking]")
    assert page.errors == []


def test_agency_sheet_sites_resolved_against_adserver(page, tmp_path):
    """Trafficking sheet d'agence : régies rapprochées des sites de l'ad server (démo hors API), choix proposé sinon."""
    from fixtures.make_agency_sheet import build
    path = str(tmp_path / "pm_agence.xlsx")
    build(path)
    _create_task(page, "PM agence")
    _analyze(page, path, validate=False)
    page.wait_for_function("document.getElementById('checks-list').textContent.includes('introuvable')", timeout=15000)
    checks = page.text_content("#checks-list")
    assert "« Facebook & Instagram » introuvable" in checks and "répertoire de démo" in checks
    assert "« TF1 + »" not in checks and "« Amnet »" not in checks                    # TF1 + -> site TF1, Amnet -> Amnet
    ids = page.evaluate("table.filter((r) => r.site_name === 'TF1').map((r) => r.site_id)")
    assert ids and all(ids)                                                          # ID du site repris
    page.click("#checks-list .site-pick button:has-text('Facebook')")               # suggestion choisie
    page.wait_for_function("!document.getElementById('checks-list').textContent.includes('Facebook & Instagram')", timeout=15000)
    names = page.evaluate("table.filter((r) => r.site_name === 'Facebook').map((r) => r.partner)")
    assert names and set(names) == {"Meta"}                                          # le nom garde le diffuseur
    page.wait_for_function("document.getElementById('preview-pane').innerText.includes('R-Notoriete-CPM-TF1-Seniors')", timeout=15000)
    assert page.errors == []


def test_reporting_home_advertiser_request_bell_and_agency_templates(page):
    """Annonceur absent : demande -> cloche de l'admin -> Ajouter -> proposé dans la liste ;
    template d'agence créé par l'admin, proposé après les 3 templates par défaut."""
    page.click(".modes [data-mode=reporting]")
    page.wait_for_function("document.querySelectorAll('#rp-home-adv option').length > 1")
    assert page.is_disabled("#rp-home-go") and not page.is_visible("#notif-count")
    page.click("#rp-req-open")
    page.fill("#rp-req-name", "Boulangerie Test")
    page.select_option("#rp-req-adserver", "adform")
    page.click("#rp-req-send")
    page.wait_for_selector("#rp-req-msg.ok")
    page.wait_for_selector("#notif-count:not([hidden])")
    page.click("#notif-btn")
    assert "Boulangerie Test" in page.text_content("#notif-panel")
    page.click("#notif-panel button:has-text('Ajouter')")
    page.wait_for_function("[...document.querySelectorAll('#rp-home-adv option')].some(o => o.textContent.includes('Boulangerie Test'))")
    assert "a été ajouté" in page.text_content("#notif-panel")
    page.mouse.click(5, 300)                                     # clic ailleurs : panneau fermé
    assert not page.is_visible("#notif-panel")
    page.click("#account-btn")
    page.click("#open-agency")
    assert "Boulangerie Test" in page.text_content("#ag-adv")
    page.fill("#ag-tpl-name", "Bilan hebdo agence")
    page.select_option("#ag-tpl-base", "check")
    page.select_option("#ag-tpl-freq", "weekly")
    page.click("#ag-tpl-form button[type=submit]")
    page.wait_for_function("document.getElementById('ag-tpl').textContent.includes('Bilan hebdo agence')")
    page.click("#agency-dialog [data-close]")
    _report_space(page, "UPSA")
    tpls = page.locator("#rp-tpls .tpl, #rp-agency-tpls .tpl").all_text_contents()
    assert "Bilan hebdo agence" in tpls[-1] and "Check de mise en ligne" in tpls[0]
    page.click("#rp-agency-tpls .tpl")
    _scope(page, "Citrate")
    page.wait_for_selector(".kpis")                              # modèle « check »
    page.click(".modes [data-mode=trafficking]")
    assert page.errors == []


def test_rename_task(page):
    """Titre : crayon (ou double-clic) -> champ, Entrée enregistre, Échap annule ; historique : crayon par tâche."""
    _report_space(page, "Banque Démo")
    page.click("#rename-task")
    page.fill("#task-title-input", "Campagne renommée")
    page.press("#task-title-input", "Escape")
    assert page.text_content("#task-title") == "Report Banque Démo"
    page.dblclick("#task-title")
    page.fill("#task-title-input", "Campagne renommée")
    page.press("#task-title-input", "Enter")
    assert page.text_content("#task-title") == "Campagne renommée" and not page.is_visible("#task-title-input")
    page.evaluate("closeTask()")
    _history(page)
    page.wait_for_selector(".task-item:has-text('Campagne renommée')")
    row = page.locator(".task-row", has=page.locator(".task-item:has-text('Campagne renommée')"))
    row.hover()
    page.once("dialog", lambda d: d.accept("Bilan banque"))
    row.locator(".ti-rename").click()
    page.wait_for_selector(".task-item:has-text('Bilan banque')")
    _open_task(page, "Bilan banque")
    page.wait_for_function("document.getElementById('task-title').textContent === 'Bilan banque'")
    page.click(".modes [data-mode=trafficking]")
    assert page.errors == []


def test_agency_memory_corrections_and_client_rules(page, tmp_path):
    """Correction faite à la main -> « Retenir » -> appliquée au PM suivant ; règles par client -> points à vérifier."""
    from fixtures.make_upsa import build
    path = tmp_path / "upsa.xlsx"
    build(str(path))
    _create_task(page, "UPSA mémoire", advertiser="UPSA")
    _analyze(page, str(path), validate=False)
    cell = page.locator('#editor tr[data-row="0"] input[data-key="format_label"]')
    before = cell.input_value()
    same = [page.locator(f'#editor tr[data-row="{i}"] input[data-key="format_label"]').input_value()
            for i in range(page.locator("#editor tbody tr").count())].count(before)
    cell.fill("Format Maison")
    cell.dispatch_event("change")
    page.wait_for_selector("#learn-bar:not([hidden])")
    assert f"« {before} » → « Format Maison »" in page.text_content("#learn-bar")
    page.select_option("#learn-scope", "UPSA")
    page.click("#learn-bar button:has-text('Retenir')")
    page.wait_for_selector("#learn-bar", state="hidden")
    # règles du client : devices autorisés + mot interdit dans les noms
    page.click("#account-btn")
    page.click("#open-agency")
    for kind, field, value in [("allowed", "device", "Tablette"), ("forbidden", "__names", "Format")]:
        page.fill("#ag-rule-adv", "UPSA")
        page.select_option("#ag-rule-kind", kind)
        page.select_option("#ag-rule-field", field)
        page.fill("#ag-rule-value", value)
        page.click("#ag-rule-form button[type=submit]")
        page.wait_for_function(f"document.getElementById('ag-rules').textContent.includes('{value}')")
    assert "Format Maison" in page.text_content("#ag-corr") and "UPSA" in page.text_content("#ag-corr")
    page.click("#agency-dialog [data-close]")
    _analyze(page, str(path), validate=False)                       # PM suivant : correction appliquée
    page.wait_for_function("document.getElementById('checks-list').textContent.includes('appliquée')")
    vals = [page.locator(f'#editor tr[data-row="{i}"] input[data-key="format_label"]').input_value()
            for i in range(page.locator("#editor tbody tr").count())]
    assert vals.count("Format Maison") == same
    checks = page.text_content("#checks-list")
    assert f"{same} correction(s) de l'agence appliquée(s)" in checks
    assert "Devices : valeur hors liste (Tablette) (règle UPSA)" in checks
    page.wait_for_function("document.getElementById('checks-list').textContent.includes('Mot interdit')")
    cell = page.locator('#editor tr[data-row="0"] input[data-key="format_label"]')   # revenir à la valeur lue : oublier ?
    cell.fill(before)
    cell.dispatch_event("change")
    page.wait_for_selector("#learn-bar:not([hidden])")
    assert "ne plus remplacer" in page.text_content("#learn-bar")
    page.click("#learn-bar button:has-text('Oublier la règle')")
    assert page.evaluate("MEM.corrections.length") == 0
    page.evaluate("Store.putSetting('client_rules', []).then(() => { MEM.rules = []; })")
    assert page.errors == []


def test_prelaunch_utm_budget_checks(page, tmp_path):
    """Avant mise en ligne : UTM incohérentes, CPM atypique, URL (vérifiées en ligne seulement)."""
    import csv
    rows = list(csv.reader(open(os.path.join(ROOT, "plan_media_template.csv"), encoding="utf-8")))
    head, body = rows[0], [r for r in rows[1:] if len(r) == len(rows[0])]
    url, budget, vol = head.index("landing_page_url"), head.index("budget_eur"), head.index("objective_impressions")
    utm = ["?utm_source=lemonde&utm_medium=display&utm_campaign=printemps", "?utm_source=lemonde&utm_medium=display&utm_campaign=Printemps",
           "?utm_source=lemonde&utm_medium=display&utm_campaign=printemps", "", "?utm_source=lemonde&utm_medium=video&utm_campaign=printemps-video",
           "?utm_source=lemonde&utm_medium=display&utm_campaign=printemps"]
    for r, q in zip(body, utm):
        r[url] += q
    body[2][budget], body[2][vol] = "30000", "100000"                  # 728x90 : CPM 300 € (médiane bien plus basse)
    body[4][budget], body[4][vol] = "5000", "400000"
    path = tmp_path / "plan_utm.csv"
    with open(path, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows([head] + body)
    _create_task(page, "Contrôles UTM")
    _analyze(page, str(path), validate=False)
    page.wait_for_function("document.getElementById('checks-list').textContent.includes('UTM')")
    checks = page.text_content("#checks-list")
    assert "UTM absentes sur 1 ligne(s)" in checks and "utm_campaign différent selon les lignes" in checks
    assert "utm_campaign écrit de plusieurs façons" in checks
    assert "CPM implicite atypique" in checks
    assert "pas en mode local" in checks                              # disponibilité des URL : Worker
    page.evaluate("closeTask()")
    assert page.errors == []


def test_campaign_monitoring_alerts_bell_and_thresholds(page):
    """Après lancement (démo) : surveillance activée depuis le check -> alertes, cloche, chiffres régie, seuils, arrêt."""
    _report_space(page, "Automobiles Démo")
    page.click(".tpl[data-tpl=check]")
    _scope(page, "Lancement SUV")
    page.wait_for_selector("#rp-watch-on")
    assert "aucune impression 24 h" in page.text_content("#rp-watch")
    page.click("#rp-watch-on")
    page.wait_for_function("document.getElementById('rp-watch').textContent.includes('Campagne surveillée')")
    alerts = page.locator(".rp-alerts li")
    assert alerts.count() >= 1 and "sans impression 24 h après le démarrage" in page.text_content(".rp-alerts")
    page.wait_for_selector("#notif-count:not([hidden])")
    page.click("#notif-btn")
    assert "Lancement SUV électrique · " in page.text_content("#notif-panel") and page.locator(".notif-item.alert").count() == alerts.count()
    page.click("#notif-btn")
    # chiffres déclarés par une régie qui diffuse : écart avec l'ad server
    page.click(".rp-regie-box summary")
    row = page.locator(".rp-regie-figs tbody tr").filter(has_not=page.locator("td:last-child:text-is('0')")).first
    site = row.locator("td").first.text_content()
    row.locator("input").first.fill("99999999")
    row.locator("input").first.dispatch_event("change")
    page.wait_for_function(f"document.querySelector('.rp-alerts') && document.querySelector('.rp-alerts').textContent.includes('{site} : écart de')")
    # seuils (admin) : affichés et enregistrés
    page.click("#account-btn")
    page.click("#open-agency")
    page.fill('#ag-th-form input[name="regie_gap"]', "15")
    page.click("#ag-th-form button[type=submit]")
    page.wait_for_selector("#ag-msg.ok")
    assert "Lancement SUV électrique" in page.text_content("#ag-monitors")
    page.click("#agency-dialog [data-close]")
    assert page.evaluate("Store.setting('alert_thresholds')")["regie_gap"] == 15
    page.click("#rp-watch-off")
    page.wait_for_selector("#rp-watch-on")
    assert page.evaluate("Store.setting('alerts').then((a) => a.filter((x) => x.status === 'open').length)") == 0
    page.click(".modes [data-mode=trafficking]")
    assert page.errors == []


def test_help_assistant_guides_without_task_data(page):
    """Assistant d'aide : conversation + nom de l'écran seulement, e-mails et URL masqués, réponse mise en forme."""
    page.evaluate("""() => {
      window.__asked = [];
      window.assistantAsk = async (messages, screen) => {
        window.__asked.push({messages: JSON.parse(JSON.stringify(messages)), screen});
        return "Voici comment faire :\\n1. Ouvrez **Reporting**\\n2. Choisissez l'annonceur\\n- puis « Espace report »";
      };
    }""")
    _report_space(page, "UPSA")
    page.click("#help-btn")
    assert page.is_visible("#help-panel") and "ne collez pas de données client" in page.text_content("#help-panel")
    assert page.locator("#help-suggest button").count() == 4
    page.fill("#help-input", "Comment envoyer à jean@agence.fr le report de https://client.fr/promo ?")
    page.click("#help-send")
    page.wait_for_selector(".help-a ol li")
    asked = page.evaluate("window.__asked")[-1]
    assert asked["screen"].startswith("Reporting — espace report")
    assert asked["messages"] == [{"role": "user", "content": "Comment envoyer à [e-mail] le report de [URL] ?"}]
    assert page.locator(".help-a ol li").count() == 2 and page.locator(".help-a strong").text_content() == "Reporting"
    assert page.locator(".help-a ul li").count() == 1
    page.fill("#help-input", "Et ensuite ?")
    page.press("#help-input", "Enter")
    page.wait_for_function("window.__asked.length === 2")
    assert [m["role"] for m in page.evaluate("window.__asked")[-1]["messages"]] == ["user", "assistant", "user"]
    page.click("#help-clear")
    assert page.locator("#help-suggest button").count() == 4
    page.keyboard.press("Escape")
    assert not page.is_visible("#help-panel")
    page.click(".modes [data-mode=trafficking]")
    assert page.errors == []
