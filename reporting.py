# -*- coding: utf-8 -*-
"""Templates de reporting Umnia (à partir des données de l'ad server).

Deux templates :
- **Check de mise en ligne** : toutes les régies et tous les emplacements de la campagne,
  avec ce qui remonte (ou pas) dans l'ad server ; petit tableau de bord et classeur Excel.
- **Report régie** : diffusion jour par jour des emplacements d'une régie, envoyé
  automatiquement à la régie selon un calendrier (aperçu et Excel identiques à l'envoi).

Données en entrée (JSON, mêmes formats pour CM360, Adform et la démo) :
- campagne : {id, name, advertiser, start, end}
- emplacements : [{id, name, site, site_id, size, start, end}]
- diffusion : [{date, placement_id, impressions, clicks}] (une ligne par jour et emplacement)
"""
import hashlib
import io
from datetime import date, timedelta
from typing import Dict, List, Optional

LIVE, NO_DATA, NOT_STARTED = "Données reçues", "Aucune donnée", "Pas encore démarré"


def _d(s: str) -> Optional[date]:
    try:
        return date.fromisoformat((s or "")[:10])
    except ValueError:
        return None


def _ctr(imp: int, clk: int) -> float:
    return round(clk / imp * 100, 2) if imp else 0.0


def _seed(*parts) -> int:
    return int(hashlib.sha1("|".join(map(str, parts)).encode()).hexdigest()[:8], 16)


# --- Démo : diffusion simulée, déterministe ------------------------------------------------

def demo_delivery(campaign: Dict, placements: List[Dict], today: str) -> List[Dict]:
    """Diffusion fictive jour par jour : environ 1 emplacement sur 4 sans donnée et une régie
    entière muette quand la campagne en compte au moins trois (cas typique d'un check)."""
    end_day = _d(today) - timedelta(days=1)
    sites = sorted({p.get("site", "") for p in placements})
    silent_site = sites[_seed(campaign.get("id"), "site") % len(sites)] if len(sites) >= 3 else None
    out = []
    for p in placements:
        s = _seed(campaign.get("id"), p.get("id"))
        if p.get("site") == silent_site or s % 4 == 0:
            continue
        start = _d(p.get("start") or campaign.get("start"))
        end = min(_d(p.get("end") or campaign.get("end")) or end_day, end_day)
        if not start or start > end:
            continue
        base = 4000 + s % 40000
        day = start
        while day <= end:
            k = _seed(p.get("id"), day.isoformat())
            imp = int(base * (0.7 + (k % 60) / 100))
            out.append({"date": day.isoformat(), "placement_id": str(p.get("id")),
                        "impressions": imp, "clicks": int(imp * (0.0008 + (k % 30) / 10000))})
            day += timedelta(days=1)
    return out


# --- Check de mise en ligne -----------------------------------------------------------------

def check(campaign: Dict, placements: List[Dict], delivery: List[Dict], today: str) -> Dict:
    """Synthèse : par emplacement (statut, volumes, 1re / dernière impression) et par régie."""
    now = _d(today)
    by_pl: Dict[str, Dict] = {}
    days = set()
    for r in delivery:
        if not (r.get("impressions") or r.get("clicks")):
            continue
        a = by_pl.setdefault(str(r["placement_id"]), {"impressions": 0, "clicks": 0, "first": "", "last": ""})
        a["impressions"] += int(r.get("impressions") or 0)
        a["clicks"] += int(r.get("clicks") or 0)
        d = r.get("date", "")
        if d:
            days.add(d)
            a["first"] = min(a["first"] or d, d)
            a["last"] = max(a["last"], d)
    rows = []
    for p in placements:
        a = by_pl.get(str(p.get("id")), {})
        start = _d(p.get("start") or campaign.get("start"))
        if a.get("impressions"):
            status = LIVE
        elif start and now and start >= now:
            status = NOT_STARTED
        else:
            status = NO_DATA
        imp, clk = a.get("impressions", 0), a.get("clicks", 0)
        rows.append({"site": p.get("site") or "Régie inconnue", "id": str(p.get("id", "")), "name": p.get("name", ""),
                     "size": p.get("size", ""), "start": p.get("start") or campaign.get("start", ""),
                     "end": p.get("end") or campaign.get("end", ""), "impressions": imp, "clicks": clk,
                     "ctr": _ctr(imp, clk), "first": a.get("first", ""), "last": a.get("last", ""), "status": status})
    rows.sort(key=lambda r: (r["site"].lower(), r["name"].lower()))
    sites: Dict[str, Dict] = {}
    for r in rows:
        s = sites.setdefault(r["site"], {"site": r["site"], "placements": 0, "live": 0, "waiting": 0,
                                         "impressions": 0, "clicks": 0})
        s["placements"] += 1
        s["live"] += r["status"] == LIVE
        s["waiting"] += r["status"] == NOT_STARTED
        s["impressions"] += r["impressions"]
        s["clicks"] += r["clicks"]
    for s in sites.values():
        s["ctr"] = _ctr(s["impressions"], s["clicks"])
        expected = s["placements"] - s["waiting"]
        s["status"] = (NOT_STARTED if not expected else LIVE if s["live"] == expected
                       else NO_DATA if not s["live"] else "Partiel")
    imp = sum(r["impressions"] for r in rows)
    clk = sum(r["clicks"] for r in rows)
    site_list = sorted(sites.values(), key=lambda s: -s["impressions"])
    return {"campaign": campaign, "today": today,
            "kpis": {"impressions": imp, "clicks": clk, "ctr": _ctr(imp, clk),
                     "placements": len(rows), "placements_live": sum(r["status"] == LIVE for r in rows),
                     "sites": len(site_list), "sites_live": sum(s["live"] > 0 for s in site_list),
                     "days": len(days), "first": min(days) if days else "", "last": max(days) if days else ""},
            "sites": site_list, "placements": rows}


# --- Report régie -------------------------------------------------------------------------

REGIE_COLUMNS = ["Date", "Emplacement", "ID emplacement", "Format", "Impressions", "Clics", "CTR (%)"]


def regie_report(campaign: Dict, placements: List[Dict], delivery: List[Dict], site: str) -> Dict:
    """Diffusion jour par jour des emplacements d'une régie : contenu exact du fichier envoyé."""
    pls = {str(p.get("id")): p for p in placements if (p.get("site") or "") == site}
    rows = []
    for r in sorted(delivery, key=lambda r: (r.get("date", ""), str(r.get("placement_id")))):
        p = pls.get(str(r.get("placement_id")))
        if not p:
            continue
        imp, clk = int(r.get("impressions") or 0), int(r.get("clicks") or 0)
        rows.append([r.get("date", ""), p.get("name", ""), str(p.get("id", "")), p.get("size", ""),
                     imp, clk, _ctr(imp, clk)])
    imp = sum(r[4] for r in rows)
    clk = sum(r[5] for r in rows)
    return {"site": site, "campaign": campaign, "columns": REGIE_COLUMNS, "rows": rows,
            "total": ["Total", "", "", "", imp, clk, _ctr(imp, clk)],
            "placements": len(pls), "placements_live": len({r[2] for r in rows if r[4]})}


# --- Classeurs Excel ----------------------------------------------------------------------

NAVY, CHAMPAGNE, GREEN, RED, GREY = "121A25", "EDE6D4", "D8EFDF", "F6D5D2", "ECE5D3"
STATUS_FILL = {LIVE: GREEN, NO_DATA: RED, NOT_STARTED: GREY, "Partiel": "F7E6C4"}


def _sheet(ws, title_lines: List[str], header: List[str], rows: List[List], widths: List[int],
           status_col: Optional[int] = None, total: Optional[List] = None) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    for i, line in enumerate(title_lines, start=1):
        ws.cell(i, 1, line).font = Font(bold=i == 1, size=13 if i == 1 else 10)
    h = len(title_lines) + 2
    for c, name in enumerate(header, start=1):
        cell = ws.cell(h, c, name)
        cell.font = Font(bold=True, color=CHAMPAGNE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r, row in enumerate(rows, start=h + 1):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(r, c, v)
            if isinstance(v, int):
                cell.number_format = "# ##0"
            elif isinstance(v, float):
                cell.number_format = "0.00"
        if status_col is not None and row[status_col] in STATUS_FILL:
            ws.cell(r, status_col + 1).fill = PatternFill("solid", fgColor=STATUS_FILL[row[status_col]])
    if total:
        r = h + 1 + len(rows)
        for c, v in enumerate(total, start=1):
            cell = ws.cell(r, c, v)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=GREY)
            if isinstance(v, int):
                cell.number_format = "# ##0"
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = ws.cell(h + 1, 1)
    ws.auto_filter.ref = f"A{h}:{chr(64 + len(header))}{h + max(len(rows), 1)}"


def _fr(iso: str) -> str:
    d = _d(iso)
    return d.strftime("%d/%m/%Y") if d else ""


def _save(wb) -> bytes:
    wb.properties.creator = "Umnia"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_check(campaign: Dict, placements: List[Dict], delivery: List[Dict], today: str) -> bytes:
    from openpyxl import Workbook
    c = check(campaign, placements, delivery, today)
    k = c["kpis"]
    head = [f"Check de mise en ligne — {campaign.get('name', '')}",
            f"Annonceur : {campaign.get('advertiser', '')} · Période : {_fr(campaign.get('start'))} → "
            f"{_fr(campaign.get('end'))} · Données au {_fr(today)}"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse"
    _sheet(ws, head, ["Indicateur", "Valeur"], [
        ["Impressions", k["impressions"]], ["Clics", k["clicks"]], ["CTR (%)", k["ctr"]],
        ["Emplacements avec données", f"{k['placements_live']} / {k['placements']}"],
        ["Régies avec données", f"{k['sites_live']} / {k['sites']}"],
        ["Première impression", _fr(k["first"])], ["Dernière impression", _fr(k["last"])]], [34, 22])
    _sheet(wb.create_sheet("Par régie"), head,
           ["Régie", "Emplacements", "Avec données", "Impressions", "Clics", "CTR (%)", "Statut"],
           [[s["site"], s["placements"], s["live"], s["impressions"], s["clicks"], s["ctr"], s["status"]]
            for s in c["sites"]], [30, 14, 14, 16, 12, 10, 20], status_col=6,
           total=["Total", k["placements"], k["placements_live"], k["impressions"], k["clicks"], k["ctr"], ""])
    _sheet(wb.create_sheet("Emplacements"), head,
           ["Régie", "Emplacement", "ID", "Format", "Début", "Fin", "Impressions", "Clics", "CTR (%)",
            "1re impression", "Dernière impression", "Statut"],
           [[r["site"], r["name"], r["id"], r["size"], _fr(r["start"]), _fr(r["end"]), r["impressions"],
             r["clicks"], r["ctr"], _fr(r["first"]), _fr(r["last"]), r["status"]] for r in c["placements"]],
           [24, 44, 12, 10, 12, 12, 14, 10, 9, 14, 16, 20], status_col=11)
    pls = {str(p.get("id")): p for p in placements}
    _sheet(wb.create_sheet("Données jour"), head,
           ["Date", "Régie", "Emplacement", "ID", "Impressions", "Clics"],
           [[_fr(r["date"]), pls.get(str(r["placement_id"]), {}).get("site", ""),
             pls.get(str(r["placement_id"]), {}).get("name", ""), str(r["placement_id"]),
             int(r.get("impressions") or 0), int(r.get("clicks") or 0)]
            for r in sorted(delivery, key=lambda r: (r.get("date", ""), str(r.get("placement_id"))))],
           [12, 24, 44, 12, 14, 10])
    return _save(wb)


def export_regie(campaign: Dict, placements: List[Dict], delivery: List[Dict], site: str, today: str) -> bytes:
    from openpyxl import Workbook
    rep = regie_report(campaign, placements, delivery, site)
    wb = Workbook()
    ws = wb.active
    ws.title = "Diffusion"
    _sheet(ws, [f"{campaign.get('advertiser', '')} — {campaign.get('name', '')} — {site}",
                f"Période : {_fr(campaign.get('start'))} → {_fr(campaign.get('end'))} · Données au {_fr(today)}"],
           rep["columns"], [[_fr(r[0])] + r[1:] for r in rep["rows"]], [12, 44, 14, 10, 14, 10, 9],
           total=rep["total"])
    return _save(wb)
