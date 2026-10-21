# -*- coding: utf-8 -*-
"""Agent de lecture de plans médias « tels qu'ils arrivent ».

Objectif : retrouver chaque information du trafficking même quand elle n'est pas
dans la colonne attendue. Étapes :

1. Lecture brute de TOUTES les feuilles (CSV ; , tab | — UTF-8 ou Windows-1252,
   XLSX, XLS) sous forme de grilles, sans supposer où est l'en-tête.
2. Repérage de la ligne d'en-tête (vocabulaire FR/EN) et de la meilleure feuille.
3. Bloc d'informations au-dessus du tableau : « Annonceur : X », « Période : du ... au
   ... », « Budget total », URL... → valeurs pour toute la campagne.
4. Colonnes : par l'en-tête (synonymes, rapprochement), PUIS par le contenu quand
   l'en-tête ne dit rien (colonne d'URL → destination, « 300x250 » → format,
   deux colonnes de dates → début/fin, « 12 500 € » → budget, balises <script> → tag,
   fichiers .mp4/.png → créatives, « CPM » → mode d'achat).
5. Ligne par ligne : cellules fusionnées complétées vers le bas, lignes de total
   ignorées, période « 02/11 au 15/12/2026 » découpée, plusieurs formats dans une
   cellule (« 300x250 / 728x90 ») → une ligne par format, format / URL / mode d'achat
   retrouvés dans le nom de l'emplacement ou une autre cellule, noms d'annonce et de
   créative générés s'ils manquent.
6. Rapport de provenance : pour chaque champ, d'où il vient.

Aucune donnée ne quitte la machine (l'option IA d'`ingest.py` reste à part).
"""
import csv
import io
import os
import re
from datetime import date, datetime, timedelta
from typing import Dict, List, Optional, Tuple

import ingest

Grid = List[List[str]]

# --- 1. Lecture brute ---------------------------------------------------------------


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _csv_grid(text: str) -> Grid:
    sample = "\n".join(text.splitlines()[:30])
    counts = {d: sample.count(d) for d in (";", ",", "\t", "|")}
    delim = max(counts, key=counts.get) if any(counts.values()) else ","
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delim)]


def read_grids(path: str) -> List[Tuple[str, Grid]]:
    """Toutes les feuilles du fichier, en grilles de chaînes."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        import xlrd
        book = xlrd.open_workbook(path)
        out = []
        for sh in book.sheets():
            grid = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate_as_datetime(cell.value, book.datemode).date().isoformat())
                    else:
                        row.append(_cell(cell.value))
                grid.append(row)
            out.append((sh.name, grid))
        return out
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        return [(ws.title, [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)])
                for ws in wb.worksheets]
    with open(path, "rb") as f:
        return [("CSV", _csv_grid(_decode(f.read())))]


# --- 1bis. Cases surlignées en jaune ---------------------------------------------------
# Si le classeur contient des cases jaunes, l'agent ne lit QUE ces cases (plus les
# libellés : ligne d'en-tête du tableau et étiquettes du bloc d'informations).

Highlights = Dict[str, set]


def is_yellow(rgb: str) -> bool:
    """Couleur « jaune » au sens large : jaune vif, jaune clair, or (teinte 40-70°)."""
    import colorsys
    try:
        r, g, b = (int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    except (ValueError, TypeError):
        return False
    h, s, v = colorsys.rgb_to_hsv(r, g, b)
    return 40 / 360 <= h <= 70 / 360 and s >= 0.15 and v >= 0.75


def _apply_tint(rgb: str, tint: float) -> str:
    out = []
    for i in (0, 2, 4):
        c = int(rgb[i:i + 2], 16)
        c = c + (255 - c) * tint if tint > 0 else c * (1 + tint)
        out.append(f"{max(0, min(255, round(c))):02X}")
    return "".join(out)


def _theme_colors(wb) -> List[str]:
    """Couleurs du thème Excel, dans l'ordre des index de cellule (lt1, dk1, lt2, dk2, accents)."""
    xml = getattr(wb, "loaded_theme", None)
    if not xml:
        return []
    if isinstance(xml, bytes):
        xml = xml.decode("utf-8", errors="replace")
    names = ["dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4",
             "accent5", "accent6", "hlink", "folHlink"]
    found = {}
    for n in names:
        m = re.search(rf"<a:{n}>.*?(?:srgbClr val=\"([0-9A-Fa-f]{{6}})\"|lastClr=\"([0-9A-Fa-f]{{6}})\")", xml, re.S)
        if m:
            found[n] = (m.group(1) or m.group(2)).upper()
    order = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4",
             "accent5", "accent6", "hlink", "folHlink"]
    return [found.get(n, "") for n in order]


def _fill_rgb(cell, theme: List[str]) -> str:
    fill = getattr(cell, "fill", None)
    if not fill or not getattr(fill, "fill_type", None) or fill.fill_type == "none":
        return ""
    col = fill.fgColor if fill.fill_type == "solid" else (fill.fgColor or fill.bgColor)
    if col is None:
        return ""
    try:
        if col.type == "rgb" and isinstance(col.rgb, str):
            return col.rgb[-6:].upper()
        if col.type == "indexed":
            from openpyxl.styles.colors import COLOR_INDEX
            return COLOR_INDEX[col.indexed][-6:].upper() if col.indexed < len(COLOR_INDEX) else ""
        if col.type == "theme" and col.theme < len(theme) and theme[col.theme]:
            return _apply_tint(theme[col.theme], col.tint or 0)
    except (AttributeError, TypeError, IndexError):
        return ""
    return ""


def read_highlights(path: str) -> Highlights:
    """{feuille: {(ligne, colonne)}} des cases surlignées en jaune (index à partir de 0).
    CSV : pas de couleurs, toujours vide."""
    ext = os.path.splitext(path)[1].lower()
    out: Highlights = {}
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)        # styles : mode complet
        theme = _theme_colors(wb)
        for ws in wb.worksheets:
            cells = set()
            for row in ws.iter_rows():
                for c in row:
                    if is_yellow(_fill_rgb(c, theme)):
                        cells.add((c.row - 1, c.column - 1))
            for rng in ws.merged_cells.ranges:          # case fusionnée jaune -> toute la plage
                if (rng.min_row - 1, rng.min_col - 1) in cells:
                    for r in range(rng.min_row - 1, rng.max_row):
                        for cc in range(rng.min_col - 1, rng.max_col):
                            cells.add((r, cc))
            if cells:
                out[ws.title] = cells
    elif ext == ".xls":
        import xlrd
        try:
            book = xlrd.open_workbook(path, formatting_info=True)
        except (NotImplementedError, xlrd.XLRDError):
            return out
        for sh in book.sheets():
            cells = set()
            for r in range(sh.nrows):
                for c in range(sh.ncols):
                    xf = book.xf_list[sh.cell_xf_index(r, c)]
                    bg = xf.background
                    rgb = book.colour_map.get(bg.pattern_colour_index) if bg.fill_pattern else None
                    if rgb and is_yellow("".join(f"{x:02X}" for x in rgb)):
                        cells.add((r, c))
            for r0, r1, c0, c1 in sh.merged_cells:
                if (r0, c0) in cells:
                    cells.update((r, c) for r in range(r0, r1) for c in range(c0, c1))
            if cells:
                out[sh.name] = cells
    return out


def _mask(grid: Grid, header_row: int, yellow: set) -> Grid:
    """Ne garde que les cases jaunes, la ligne d'en-tête et, au-dessus du tableau,
    les étiquettes (« Annonceur », « Période »...) suivies d'une case jaune."""
    out = []
    for r, row in enumerate(grid):
        if r == header_row:
            out.append(list(row))
            continue
        new = [v if (r, c) in yellow else "" for c, v in enumerate(row)]
        if r < header_row:
            for c, v in enumerate(row):
                if v and not new[c] and _context_key(v) and any((r, x) in yellow for x in range(c + 1, len(row))):
                    new[c] = v
        out.append(new)
    return out


# --- Reconnaissance des valeurs ----------------------------------------------------

RE_URL = re.compile(r"(https?://[^\s\"'<>]+|www\.[^\s\"'<>]+\.[a-z]{2,}[^\s\"'<>]*)", re.I)
RE_SIZE = re.compile(r"(?<!\d)(\d{2,4})\s*[xX×*]\s*(\d{2,4})(?!\d)")
RE_TAG = re.compile(r"<\s*(script|iframe|ins|img|noscript)\b|javascript:", re.I)
RE_FILE = re.compile(r"\.(png|jpe?g|gif|wbmp|webp|zip|adz|mp4|mov|m4v|webm|mpe?g|avi|wmv|mp3|ogg|oga|ogv)\s*$", re.I)
RE_PRICING = re.compile(r"\b(cpm|cpc|cpa|cpv|cpcv|vcpm|flat|forfait|co[uû]t pour mille)\b", re.I)
RE_TOTAL = re.compile(r"^\s*(sous[\s_-]*)?total\b|^\s*total\s", re.I)
RE_MONEY = re.compile(r"(€|\beur\b|\beuros?\b|k€)", re.I)
RE_DATE_TOKEN = re.compile(
    r"\d{4}-\d{2}-\d{2}|\d{1,2}[/.\-]\d{1,2}(?:[/.\-]\d{2,4})?")
VIDEO_HINT = re.compile(r"pre[\s_-]?roll|in[\s_-]?stream|vid[eé]o|outstream|mid[\s_-]?roll", re.I)


def parse_date(s: str, default_year: Optional[int] = None) -> str:
    """Date -> AAAA-MM-JJ, sinon ''. Formats européens (JJ/MM), ISO, AA sur 2 chiffres,
    numéro de série Excel, et JJ/MM sans année (année fournie)."""
    s = (s or "").strip()
    if not s:
        return ""
    head = re.split(r"[ T]", s)[0]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%Y", "%d/%m/%y", "%d.%m.%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(head, fmt).date().isoformat()
        except ValueError:
            pass
    m = re.fullmatch(r"(\d{1,2})[/.\-](\d{1,2})", head)
    if m and default_year:
        try:
            return date(default_year, int(m.group(2)), int(m.group(1))).isoformat()
        except ValueError:
            return ""
    if re.fullmatch(r"\d{5}(\.0)?", head):                      # série Excel (1900)
        n = int(float(head))
        if 30000 < n < 80000:
            return (date(1899, 12, 30) + timedelta(days=n)).isoformat()
    return ""


def parse_period(s: str) -> Tuple[str, str]:
    """« du 02/11/2026 au 15/12/2026 », « 02/11 - 15/12/2026 », « 2026-11-02 → 2026-12-15 »."""
    toks = RE_DATE_TOKEN.findall(s or "")
    if len(toks) < 2:
        return "", ""
    year = None
    for t in reversed(toks):
        m = re.search(r"(\d{4})$|^(\d{4})", t)
        if m:
            year = int(m.group(1) or m.group(2))
            break
    a, b = parse_date(toks[0], year), parse_date(toks[1], year)
    if a and b and a > b and re.fullmatch(r"\d{1,2}[/.\-]\d{1,2}", toks[0]) and year:
        a = parse_date(toks[0], year - 1)                         # « 15/12 au 10/01/2027 »
    return a, b


def _kind(v: str) -> str:
    """Nature d'une valeur : url, size, sizes, tag, file, pricing, period, date, money, int, text, ''."""
    v = (v or "").strip()
    if not v:
        return ""
    if RE_TAG.search(v):
        return "tag"
    if RE_URL.fullmatch(v):
        return "url"
    if RE_FILE.search(v) and " " not in v.strip():
        return "file"
    sizes = RE_SIZE.findall(v)
    if sizes and len(re.sub(r"[\d\sxX×*/,;+\-px|et]", "", v, flags=re.I)) <= 2:
        return "sizes" if len(sizes) > 1 else "size"
    if RE_PRICING.fullmatch(v.strip()) or (RE_PRICING.search(v) and len(v) <= 20):
        return "pricing"
    if len(RE_DATE_TOKEN.findall(v)) >= 2 and parse_period(v)[0]:
        return "period"
    if parse_date(v):
        return "date"
    num = re.sub(r"[\s \xa0€]|eur(os?)?|k", "", v, flags=re.I).replace(",", ".")
    if re.fullmatch(r"-?\d+(\.\d+)?", num):
        return "money" if RE_MONEY.search(v) else "int"
    return "text"


# Noms de formats usuels (FR/EN, IAB) → dimensions. Ordre : du plus précis au plus général.
FORMAT_NAMES = [
    (r"grand[\s_-]*angle|half[\s_-]*page|demi[\s_-]*page|large[\s_-]*pav", "300x600"),
    (r"billboard|masthead|maxi[\s_-]*banni|super[\s_-]*banni", "970x250"),
    (r"large[\s_-]*leaderboard|super[\s_-]*leaderboard", "970x90"),
    (r"m[eé]ga[\s_-]*banni|leaderboard|banni[eè]re[\s_-]*haute|top[\s_-]*banner", "728x90"),
    (r"wide[\s_-]*sky|large[\s_-]*sky|sky[\s_-]*large", "160x600"),
    (r"skyscraper|\bsky\b|gratte[\s_-]*ciel", "120x600"),
    (r"interstitiel|interstitial|plein[\s_-]*[eé]cran", "320x480"),
    (r"large[\s_-]*mobile|mobile[\s_-]*large|banni[eè]re[\s_-]*mobile[\s_-]*large", "320x100"),
    (r"banni[eè]re[\s_-]*mobile|mobile[\s_-]*banner|\bbanner[\s_-]*mobile|smartphone[\s_-]*banni", "320x50"),
    (r"pav[eé]|\bmpu\b|medium[\s_-]*rectangle|rectangle|carr[eé]", "300x250"),
    (r"pre[\s_-]?roll|mid[\s_-]?roll|in[\s_-]?stream|vid[eé]o|spot", "640x360"),
]
TRACKING_SIZE = "1x1"
VIDEO_SIZE = "640x360"


def category_from_name(text: str) -> str:
    """Catégorie d'emplacement déduite d'un nom de format : vidéo ou display (IAB)."""
    return "video" if VIDEO_HINT.search(text or "") else ("display" if size_from_name(text) else "")


def size_from_name(text: str) -> str:
    """« Pavé », « Méga-bannière », « Preroll 15s »… → dimensions ; '' si inconnu."""
    t = text or ""
    for pattern, size in FORMAT_NAMES:
        if re.search(pattern, t, re.I):
            return size
    return ""


# --- 2. En-tête et feuille -------------------------------------------------------------

def _header_score(row: List[str]) -> int:
    norm = [ingest._norm(c) for c in row if c and _kind(c) in ("text", "")]
    if len(norm) < 2:
        return 0
    mapping, _ = ingest.map_headers(norm)
    return len(set(mapping.values()))


def find_table(grids: List[Tuple[str, Grid]]) -> Tuple[str, Grid, int]:
    """(feuille, grille, index de la ligne d'en-tête). Choix : meilleure ligne
    d'en-tête (vocabulaire reconnu), puis nombre de lignes de données."""
    best = None
    for name, grid in grids:
        for i, row in enumerate(grid[:40]):
            s = _header_score(row)
            if s >= 2:
                data = sum(1 for r in grid[i + 1:] if sum(1 for c in r if c) >= 2)
                key = (s, data)
                if best is None or key > best[0]:
                    best = (key, name, grid, i)
    if best:
        return best[1], best[2], best[3]
    # Pas d'en-tête reconnu : feuille la plus remplie, première ligne pleine
    name, grid = max(grids, key=lambda g: sum(1 for r in g[1] if any(r)))
    first = next((i for i, r in enumerate(grid) if sum(1 for c in r if c) >= 2), 0)
    return name, grid, first


# --- 3. Bloc d'informations au-dessus du tableau ---------------------------------------

CONTEXT_KEYS = {
    "advertiser_name": ["annonceur", "client", "advertiser", "marque", "brand"],
    "campaign_name": ["campagne", "campaign", "operation", "nom_de_la_campagne", "projet"],
    "period": ["periode", "dates", "date", "flight", "vague", "periode_de_diffusion"],
    "campaign_start": ["debut", "date_de_debut", "start", "start_date"],
    "campaign_end": ["fin", "date_de_fin", "end", "end_date"],
    "budget_eur": ["budget", "budget_total", "budget_net", "budget_ht", "montant"],
    "landing_page_url": ["url", "url_de_destination", "landing", "lp", "site_annonceur"],
}


def _context_key(label: str) -> Optional[str]:
    n = ingest._norm(label)
    for key, syns in CONTEXT_KEYS.items():
        if n in syns or any(n.startswith(s + "_") for s in syns):
            return key
    return None


def read_context(cells: List[List[str]]) -> Dict[str, Tuple[str, str]]:
    """{champ: (valeur, cellule)} trouvés dans les lignes au-dessus du tableau."""
    found: Dict[str, Tuple[str, str]] = {}
    for r, row in enumerate(cells):
        for c, v in enumerate(row):
            if not v:
                continue
            where = f"ligne {r + 1}, colonne {c + 1}"
            label, value = None, None
            m = re.match(r"^\s*([^:]{2,40}?)\s*:\s*(.+)$", v)
            if m and not v.lower().startswith("http"):
                label, value = m.group(1), m.group(2)
            elif _context_key(v):
                nxt = next((x for x in row[c + 1:] if x), "")
                label, value = v, nxt
            if label and value:
                key = _context_key(label)
                if key and key not in found:
                    found[key] = (value.strip(), where)
            if "landing_page_url" not in found and RE_URL.fullmatch(v):
                found["landing_page_url"] = (v, where)
    if "period" in found:
        a, b = parse_period(found["period"][0])
        if a and b:
            found.setdefault("campaign_start", (a, found["period"][1]))
            found.setdefault("campaign_end", (b, found["period"][1]))
    for k in ("campaign_start", "campaign_end"):
        if k in found:
            d = parse_date(found[k][0])
            if d:
                found[k] = (d, found[k][1])
            else:
                del found[k]
    return found


# --- 4. Colonnes -------------------------------------------------------------------------

CONTENT_TARGET = {"url": "landing_page_url", "size": "size", "sizes": "size", "tag": "snippet",
                  "file": "creative_path", "pricing": "pricing_type", "money": "budget_eur"}
FILL_DOWN = ["advertiser_name", "campaign_name", "site_name", "site_id", "campaign_start",
             "campaign_end", "pricing_type", "landing_page_url"]


def _profile(values: List[str]) -> Tuple[str, float]:
    kinds = [_kind(v) for v in values if v]
    if not kinds:
        return "", 0.0
    best = max(set(kinds), key=kinds.count)
    share = kinds.count(best) / len(kinds)
    if best == "size" and "sizes" in kinds:
        share = (kinds.count("size") + kinds.count("sizes")) / len(kinds)
    return best, share


def map_columns(headers: List[str], data: Grid) -> Tuple[Dict[int, str], List[str], List[str]]:
    """{index de colonne: champ}, provenance, colonnes ignorées."""
    norm = [ingest._norm(h) for h in headers]
    by_header, _ = ingest.map_headers([n for n in norm if n])
    cols: Dict[int, str] = {}
    prov: List[str] = []
    for i, n in enumerate(norm):
        if n and n in by_header and by_header[n] not in cols.values():
            cols[i] = by_header[n]
            prov.append(f"{by_header[n]} ← colonne « {headers[i]} »")
    # « Objectif » : texte (notoriété, trafic...) = objectif média ; chiffres = volume
    for i, f in list(cols.items()):
        if f in ("objective", "objective_impressions"):
            kind, share = _profile([r[i] for r in data if i < len(r) and r[i]])
            want = "objective_impressions" if kind in ("int", "money") else "objective"
            if want != f and want not in cols.values():
                cols[i] = want
    # Période dans une seule colonne (en-tête « Période », « Dates »...)
    for i, n in enumerate(norm):
        if i in cols or not n:
            continue
        values = [r[i] for r in data if i < len(r) and r[i]]
        kind, share = _profile(values)
        if kind == "period" and share >= .6 and "campaign_start" not in cols.values():
            cols[i] = "period"
            prov.append(f"campaign_start + campaign_end ← période découpée, colonne « {headers[i]} »")
    # Par le contenu, pour les colonnes non reconnues
    date_cols = []
    for i in range(len(headers)):
        if i in cols:
            continue
        values = [r[i] for r in data if i < len(r) and r[i]]
        kind, share = _profile(values)
        if share < .6 or not values:
            continue
        target = CONTENT_TARGET.get(kind)
        if kind == "date":
            date_cols.append(i)
        elif kind == "int" and re.search(r"imp|vol|kpi|objectif|qt|quantit", norm[i] or ""):
            target = "objective_impressions"
        if target and target not in cols.values():
            cols[i] = target
            prov.append(f"{target} ← contenu de la colonne « {headers[i] or f'n°{i + 1}'} » "
                        f"({share:.0%} de valeurs de type {kind})")
    taken = set(cols.values())
    free_dates = [i for i in date_cols if i not in cols]
    if free_dates and "campaign_start" not in taken:
        med = lambda i: sorted(parse_date(r[i]) for r in data if i < len(r) and parse_date(r[i]))[0]
        free_dates.sort(key=med)
        cols[free_dates[0]] = "campaign_start"
        prov.append(f"campaign_start ← dates de la colonne « {headers[free_dates[0]] or free_dates[0] + 1} »")
        if len(free_dates) > 1 and "campaign_end" not in taken:
            cols[free_dates[1]] = "campaign_end"
            prov.append(f"campaign_end ← dates de la colonne « {headers[free_dates[1]] or free_dates[1] + 1} »")
    # Nom d'emplacement manquant : colonne de texte la plus variée
    if "placement_name" not in cols.values():
        cands = []
        for i in range(len(headers)):
            if i in cols:
                continue
            values = [r[i] for r in data if i < len(r) and r[i]]
            kind, share = _profile(values)
            if kind == "text" and share >= .8:
                cands.append((len(set(values)), i))
        if cands:
            i = max(cands)[1]
            cols[i] = "placement_name"
            prov.append(f"placement_name ← colonne de libellés « {headers[i] or i + 1} » (la plus détaillée)")
    ignored = [h or f"colonne {i + 1}" for i, h in enumerate(headers) if i not in cols and h]
    return cols, prov, ignored


# --- 5. Lignes ---------------------------------------------------------------------------

def _is_total(row: List[str]) -> bool:
    return any(RE_TOTAL.match(c or "") for c in row[:4])


def extract(grids: List[Tuple[str, Grid]], campaign_name: str = "",
            advertiser_name: str = "", highlights: Optional[Highlights] = None
            ) -> Tuple[List[Dict[str, str]], List[str]]:
    """Lignes canoniques + rapport de provenance. Avec `highlights` (cases jaunes),
    seules les cases surlignées sont lues."""
    yellow_mode = bool(highlights and any(highlights.values()))
    if yellow_mode:
        sheet, grid, h = find_table([(n, g) for n, g in grids if highlights.get(n)])
    else:
        sheet, grid, h = find_table(grids)
    width = max((len(r) for r in grid), default=0)
    grid = [r + [""] * (width - len(r)) for r in grid]
    if yellow_mode:
        grid = _mask(grid, h, highlights[sheet])
    headers = grid[h] if h < len(grid) else []
    context = read_context(grid[:h])
    # Données : jusqu'à 3 lignes vides consécutives (en mode jaune : tout le tableau,
    # les lignes non surlignées étant simplement vides)
    data, blanks = [], 0
    for row in grid[h + 1:]:
        if not any(row):
            blanks += 1
            if blanks >= 3 and data and not yellow_mode:
                break
            continue
        blanks = 0
        data.append(row)
    report = []
    if yellow_mode:
        n = sum(len(v) for v in highlights.values())
        report.append(f"Mode cases jaunes : seules les {n} case(s) surlignée(s) en jaune sont lues "
                      "(plus les libellés d'en-tête), le reste du fichier est ignoré.")
    report.append(f"Feuille « {sheet} », en-tête ligne {h + 1}, {len(data)} ligne(s) de données.")
    cols, prov, ignored = map_columns(headers, data)
    report += prov
    for k, (v, where) in context.items():
        if k != "period":
            report.append(f"{k} ← bloc d'en-tête du plan ({where}) : {v}")

    rows: List[Dict[str, str]] = []
    skipped = 0
    last: Dict[str, str] = {}
    filled, derived = set(), set()
    for raw in data:
        if _is_total(raw):
            skipped += 1
            continue
        r = {f: raw[i] for i, f in cols.items() if f != "period"}
        for i, f in cols.items():
            if f == "period":
                a, b = parse_period(raw[i])
                r["campaign_start"], r["campaign_end"] = a, b
        others = [raw[i] for i in range(width) if i not in cols and raw[i]]
        names = [r.get(k, "") for k in ("placement_name", "creative_name", "ad_name")]
        # Cellules fusionnées / valeurs écrites une seule fois
        for f in FILL_DOWN:
            if not r.get(f) and last.get(f) and f in cols.values():
                r[f] = last[f]
                filled.add(f)
        # Informations au mauvais endroit : on les cherche dans la ligne
        category = ""
        written_label = r.get("size", "")
        if RE_SIZE.search(r.get("size", "")):
            category = "display"
        else:
            written = r.get("size", "")
            r["size"] = ""
            for t in names + others:
                m = RE_SIZE.search(t or "")
                if m:
                    r["size"] = f"{m.group(1)}x{m.group(2)}"
                    category = "display"
                    derived.add("size ← retrouvé dans le nom de l'emplacement ou une autre cellule")
                    break
            if not r["size"]:
                for t in [written] + names + others:
                    sz, cat = size_from_name(t), category_from_name(t)
                    if sz:
                        r["size"], category = sz, cat
                        derived.add("size ← nom de format IAB converti en dimensions, emplacement "
                                    "display (Pavé → 300x250, Méga-bannière → 728x90…)"
                                    if cat == "display" else
                                    "size ← format vidéo (preroll, in-stream…) : emplacement vidéo in-stream")
                        break
            if not r["size"]:
                r["size"], category = TRACKING_SIZE, "tracking"
                derived.add(f"size ← aucun format reconnaissable : {TRACKING_SIZE}, "
                            "emplacement de tracking")
        if VIDEO_HINT.search(" ".join(names + [r.get("size", "")])) or \
                re.search(r"\.(mp4|mov|m4v|webm)\s*$", r.get("creative_path", ""), re.I):
            category = "video"
        if not r.get("placement_category"):
            r["placement_category"] = category
        if not r.get("format_label"):
            r["format_label"] = written_label or r.get("size", "")
        if not r.get("landing_page_url"):
            for t in others:
                m = RE_URL.search(t)
                if m:
                    r["landing_page_url"] = m.group(1)
                    derived.add("landing_page_url ← URL trouvée dans une autre cellule de la ligne")
                    break
        if not r.get("pricing_type"):
            for t in others + names:
                m = RE_PRICING.search(t or "")
                if m:
                    r["pricing_type"] = m.group(1)
                    derived.add("pricing_type ← mode d'achat trouvé dans une autre cellule")
                    break
        if not r.get("snippet"):
            for t in others:
                if RE_TAG.search(t):
                    r["snippet"] = t
                    derived.add("snippet ← balise trouvée dans une autre cellule")
                    break
        # Valeurs de campagne : bloc d'en-tête puis paramètres
        for f in ("advertiser_name", "campaign_name", "campaign_start", "campaign_end", "landing_page_url"):
            if not r.get(f) and f in context:
                r[f] = context[f][0]
        if not r.get("campaign_name") and campaign_name:
            r["campaign_name"] = campaign_name
        if not r.get("advertiser_name") and advertiser_name:
            r["advertiser_name"] = advertiser_name
        if not any(r.get(k) for k in ("placement_name", "size", "ad_name", "creative_name")):
            skipped += 1                                          # ligne de titre ou de note
            continue
        for f in FILL_DOWN:
            if r.get(f):
                last[f] = r[f]
        # Plusieurs formats dans une cellule -> une ligne par format
        sizes = [f"{a}x{b}" for a, b in RE_SIZE.findall(r.get("size", ""))] or [r.get("size", "")]
        for sz in sizes:
            line = dict(r, size=sz)
            if len(sizes) > 1:
                line["format_label"] = sz
                base = line.get("placement_name") or "Emplacement"
                if sz not in base:
                    line["placement_name"] = f"{base}_{sz}"
                derived.add("plusieurs formats dans une cellule → une ligne par format")
            rows.append(line)

    # Noms générés et normalisation
    gen = set()
    seen_creas: set = set()
    for line in rows:
        if not line.get("placement_name"):
            line["placement_name"] = " ".join(x for x in (line.get("site_name"), line.get("size")) if x) or "Emplacement"
            gen.add("placement_name")
        if not line.get("ad_name"):
            line["ad_name"] = line["placement_name"]
            gen.add("ad_name")
        if not line.get("creative_name"):
            base = f"{line['ad_name']}_{line['message']}" if line.get("message") else f"{line['ad_name']}_crea"
            name, n = base, 2
            while (line["ad_name"], name) in seen_creas:      # plusieurs créatives sur la même annonce
                name, n = f"{base}_{n}", n + 1
            line["creative_name"] = name
            gen.add("creative_name")
        seen_creas.add((line["ad_name"], line["creative_name"]))
        if not line.get("creative_category") and not line.get("creative_path") and not line.get("snippet") \
                and VIDEO_HINT.search(line.get("placement_name", "")):
            pass   # sans fichier : créative de tracking (règle Trafficker), même sur un emplacement vidéo
        for f, fn in ingest.NORMALIZERS.items():
            if line.get(f):
                line[f] = fn(str(line[f]))
        for f in ("campaign_start", "campaign_end"):
            if line.get(f):
                line[f] = parse_date(line[f]) or line[f]
    _campaign_period(rows, report)
    _aggregate_campaign_totals(rows, cols, report)

    report += sorted(derived)
    if filled:
        report.append("Cellules fusionnées complétées vers le bas : " + ", ".join(sorted(filled)))
    if gen:
        report.append("Noms générés (absents du plan) : " + ", ".join(sorted(gen)))
    if skipped:
        report.append(f"{skipped} ligne(s) de total, de titre ou de note ignorée(s).")
    if ignored:
        report.append("Colonnes non utilisées : " + ", ".join(ignored))
    missing = [f for f in ingest.REQUIRED if not any(x.get(f) for x in rows)]
    if missing:
        report.append("Introuvables dans le plan : " + ", ".join(missing)
                      + (" — renseignez la campagne dans les réglages." if "campaign_name" in missing else ""))
    return rows, report


def _campaign_period(rows: List[Dict[str, str]], report: List[str]) -> None:
    """Dates propres à chaque ligne (vagues, emplacements décalés) : la campagne
    couvre du premier début à la dernière fin."""
    by_camp: Dict[str, Tuple[List[str], List[str]]] = {}
    for line in rows:
        s, e = by_camp.setdefault(line.get("campaign_name", ""), ([], []))
        if line.get("campaign_start"):
            s.append(line["campaign_start"])
        if line.get("campaign_end"):
            e.append(line["campaign_end"])
    for camp, (s, e) in by_camp.items():
        if len(set(s)) > 1 or len(set(e)) > 1:
            lo, hi = min(s), max(e)
            for line in rows:
                if line.get("campaign_name", "") == camp:
                    line["campaign_start"], line["campaign_end"] = lo, hi
            report.append(f"Dates différentes selon les lignes : la campagne « {camp} » "
                          f"couvre du {lo} au {hi}.")


def _aggregate_campaign_totals(rows: List[Dict[str, str]], cols: Dict[int, str],
                               report: List[str]) -> None:
    """Budget et impressions lus par ligne → total de la campagne (le plan Trafficker
    porte un budget et un objectif par campagne)."""
    for f in ("budget_eur", "objective_impressions"):
        if f not in cols.values():
            continue
        totals: Dict[str, float] = {}
        for line in rows:
            try:
                totals[line.get("campaign_name", "")] = totals.get(line.get("campaign_name", ""), 0) \
                    + float(line.get(f) or 0)
            except ValueError:
                pass
        for line in rows:
            t = totals.get(line.get("campaign_name", ""))
            if t:
                line[f] = f"{t:g}" if f == "budget_eur" else str(int(t))
        if totals:
            report.append(f"{f} : somme des lignes par campagne.")


def is_canonical(grids: List[Tuple[str, Grid]]) -> bool:
    """Plan déjà au format Trafficker (en-tête canonique complet en ligne 1)."""
    if not grids or not grids[0][1]:
        return False
    head = [ingest._norm(c) for c in grids[0][1][0]]
    return all(f in head for f in ingest.REQUIRED)
